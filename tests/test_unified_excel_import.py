import hashlib
import json
import sqlite3
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker

from app.cli.import_excel import build_parser
from app.core.database import Base
from app.importers import unified_excel_import as import_orchestrator
from app.importers.unified_excel_import import ImportPostCommitError, ImportPreflightError, apply_unified_import
from app.importers.unified_excel_preview import MATERIAL_HEADERS, WORK_HEADERS, run_preview
from app.models.models import (
    Equipment, Material, Project, StatusEnum, WorkEquipmentLink,
    WorkItem, WorkMaterialLink,
)
from app.repositories import unified_excel_import as import_repository
from app.repositories.project_budget_summaries import get_project_budget_summary
from app.repositories.unified_excel_import import ImportPersistenceError
from app.schemas.import_execution import IMPORT_CONFIRMATION


PROJECT_ID = "PRJ-ALTAI-R7-B"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def material_id(index: int) -> str:
    prefix = "PRD-ALTAI-B" if index <= 138 else "PRD-ALTAI-6R"
    return f"{prefix}-{index:04d}"


def make_workbook(path: Path) -> Decimal:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ("Заавар", "Материалын үнэ", "Ажил-Тээвэр-Механизм", "Нэгтгэлийн хяналт"):
        workbook.create_sheet(name)
    material_sheet = workbook["Материалын үнэ"]
    work_sheet = workbook["Ажил-Тээвэр-Механизм"]
    for column, value in enumerate(MATERIAL_HEADERS, 1):
        material_sheet.cell(4, column, value)
    for column, value in enumerate(WORK_HEADERS, 1):
        work_sheet.cell(4, column, value)
    special_names = {
        1: "Unique A", 811: "Unique A", 4: "Unique B", 812: "Unique B",
        2: "Ambiguous A", 3: "Ambiguous A", 813: "Ambiguous A", 814: "Ambiguous A",
        5: "Ambiguous B", 6: "Ambiguous B", 815: "Ambiguous B", 816: "Ambiguous B",
    }
    for index in range(1, 987):
        material_sheet.append([
            material_id(index), "Category", special_names.get(index, f"Material {index}"),
            "м²", None, float(index * 10) if index <= 810 else None,
            None, None, None, None,
        ])
    labor = Decimal("0")
    for index in range(1, 66):
        quantity = Decimal(index)
        rate = Decimal("1000")
        labor += (quantity * rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        work_sheet.append([
            f"WRK-ALTAI-B-{index:03d}", "Category", f"Work {index}",
            float(quantity), "м²", float(rate), None, None, None, None,
            None, None, None, None, None,
        ])
    workbook.save(path)
    return labor


def make_database(path: Path) -> None:
    engine = create_engine("sqlite:///" + path.as_posix())
    Base.metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE alembic_version (version_num VARCHAR(32) NOT NULL PRIMARY KEY)"))
        connection.execute(text("INSERT INTO alembic_version VALUES ('0004_work_masters')"))
    session = sessionmaker(bind=engine)()
    project = Project(project_id=PROJECT_ID, name="Project")
    session.add(project); session.flush()
    work = WorkItem(
        work_id=f"{PROJECT_ID}-WRK-001", project_id=project.id, name="Work 1",
        unit="м²", quantity=1, labor_unit_rate=1000, labor_total=1000,
        status=StatusEnum.ACTIVE,
    )
    session.add(work); session.flush()
    existing = []
    for index in range(981, 987):
        material = Material(
            material_id=material_id(index), name=f"Material {index}",
            normalized_unit="м²", unit_price=float(index), status=StatusEnum.ACTIVE,
        )
        session.add(material); existing.append(material)
    session.flush()
    for material in existing:
        session.add(WorkMaterialLink(
            work_id=work.id, material_id=material.id, consumption_rate=1,
            waste_percentage=0, calculated_quantity=1, status=StatusEnum.ACTIVE,
        ))
    equipment = Equipment(equipment_id="E-1", type="Crane", unit_rate=100)
    session.add(equipment); session.flush()
    session.add(WorkEquipmentLink(
        work_item_id=work.id, equipment_id=equipment.id, usage_quantity=1,
        agreed_unit_rate=100, status=StatusEnum.ACTIVE,
    ))
    session.commit(); session.close(); engine.dispose()


@pytest.fixture()
def import_case(tmp_path: Path):
    for name in ("input", "database", "plans", "backups", "reports", "repository"):
        (tmp_path / name).mkdir()
    workbook = tmp_path / "input" / "unified.xlsx"
    database = tmp_path / "database" / "hubstar.db"
    expected_labor = make_workbook(workbook)
    make_database(database)
    plan_dir = tmp_path / "plans" / "preview"
    preview = run_preview(
        file_path=workbook, project_id=PROJECT_ID, database_path=database,
        report_dir=plan_dir, repository_root=tmp_path / "repository",
        generated_at_utc="2026-01-01T00:00:00Z",
    )
    return {
        "root": tmp_path, "workbook": workbook, "database": database,
        "plan": Path(preview.report_paths[0]), "expected_labor": expected_labor,
    }


def apply(case, suffix="one"):
    backup = case["root"] / "backups" / f"backup-{suffix}.db"
    report = case["root"] / "reports" / f"execution-{suffix}"
    result = apply_unified_import(
        file_path=case["workbook"], project_id=PROJECT_ID,
        database_path=case["database"], plan_path=case["plan"],
        expected_workbook_sha256=sha(case["workbook"]),
        expected_database_sha256=sha(case["database"]),
        backup_path=backup, report_dir=report, confirm=IMPORT_CONFIRMATION,
        repository_root=case["root"] / "repository",
    )
    return result, backup, report


def test_fixture_scale_apply_backup_transaction_and_reports(import_case):
    pre_hash = sha(import_case["database"])
    result, backup, report = apply(import_case)
    assert result.summary["status"] == "SUCCESS"
    assert result.summary["created"] == {"work_masters": 65, "work_items": 64, "materials": 980}
    assert result.summary["linked"] == 1
    assert result.summary["skipped"] == {"work_masters": 0, "work_items": 1, "materials": 6}
    assert backup.exists() and sha(backup) == result.summary["backup_sha256"]
    assert len(result.report_paths) == 5
    assert all(Path(path).read_bytes().startswith(b"\xef\xbb\xbf") for path in result.report_paths[1:])
    assert str(import_case["root"]) not in json.dumps(result.summary)
    with sqlite3.connect(backup) as backup_db:
        assert backup_db.execute("pragma quick_check").fetchone()[0] == "ok"
        assert backup_db.execute("pragma foreign_key_check").fetchall() == []
        assert backup_db.execute("select count(*) from work_items").fetchone()[0] == 1
        assert backup_db.execute("select count(*) from materials").fetchone()[0] == 6

    connection = sqlite3.connect(import_case["database"])
    counts = {
        table: connection.execute(f"select count(*) from {table}").fetchone()[0]
        for table in ("work_masters", "work_items", "materials", "work_material_links",
                      "equipments", "work_equipment_links", "projects", "transports",
                      "import_batches", "audit_logs")
    }
    assert counts == {
        "work_masters": 65, "work_items": 65, "materials": 986,
        "work_material_links": 6, "equipments": 1, "work_equipment_links": 1,
        "projects": 1, "transports": 0, "import_batches": 1, "audit_logs": 1,
    }
    assert connection.execute("select count(*) from materials where status='ACTIVE'").fetchone()[0] == 818
    assert connection.execute("select count(*) from materials where status='NEEDS_REVIEW'").fetchone()[0] == 168
    assert connection.execute("select count(*) from materials where unit_price is null").fetchone()[0] == 168
    assert connection.execute("select unit_price,status from materials where material_id=?", (material_id(811),)).fetchone() == (10.0, "ACTIVE")
    assert connection.execute("select unit_price,status from materials where material_id=?", (material_id(812),)).fetchone() == (40.0, "ACTIVE")
    assert connection.execute("select name,unit_price,status from materials where material_id=?", (material_id(981),)).fetchone() == ("Material 981", 981.0, "ACTIVE")
    facade = connection.execute(
        "select name,unit,quantity,labor_unit_rate,labor_total,status,work_master_ref_id "
        "from work_items where work_id=?", (f"{PROJECT_ID}-WRK-001",)
    ).fetchone()
    assert facade[:6] == ("Work 1", "м²", 1.0, 1000.0, 1000.0, "ACTIVE")
    assert facade[6] is not None
    labor = Decimal(str(connection.execute("select sum(labor_total) from work_items").fetchone()[0]))
    assert labor == import_case["expected_labor"]
    assert connection.execute(
        "select count(*) from work_items wi left join work_material_links l on l.work_id=wi.id "
        "where l.id is null"
    ).fetchone()[0] == 64
    connection.close()
    verify_engine = create_engine("sqlite:///" + import_case["database"].as_posix())
    verify_session = sessionmaker(bind=verify_engine)()
    budget = get_project_budget_summary(verify_session, PROJECT_ID)
    assert budget["work_item_count"] == 65
    assert budget["no_materials_work_count"] == 64
    assert budget["pricing_status"] == "INCOMPLETE"
    assert Decimal(str(budget["labor_subtotal_known"])) == import_case["expected_labor"]
    verify_session.close(); verify_engine.dispose()
    assert pre_hash != sha(import_case["database"])
    assert {path.name for path in report.iterdir()} == {
        "import-execution-summary.json", "work-master-import-result.csv",
        "work-item-import-result.csv", "material-import-result.csv",
        "import-execution-conflicts.csv",
    }


def test_idempotent_second_run_is_no_changes_without_backup_or_database_write(import_case):
    apply(import_case)
    second_plan_dir = import_case["root"] / "plans" / "second"
    preview = run_preview(
        file_path=import_case["workbook"], project_id=PROJECT_ID,
        database_path=import_case["database"], report_dir=second_plan_dir,
        repository_root=import_case["root"] / "repository",
    )
    import_case["plan"] = Path(preview.report_paths[0])
    before = sha(import_case["database"])
    result, backup, _ = apply(import_case, "two")
    assert result.summary["status"] == "NO_CHANGES"
    assert result.summary["database_writes_performed"] is False
    assert not backup.exists()
    assert sha(import_case["database"]) == before


@pytest.mark.parametrize("checkpoint", ["before_domain", "after_masters", "after_work_items", "before_commit"])
def test_failure_at_transaction_checkpoints_rolls_back_domain_and_metadata(import_case, monkeypatch, checkpoint):
    before = sha(import_case["database"])
    def fail(name):
        if name == checkpoint:
            raise ImportPersistenceError("forced failure")
    monkeypatch.setattr(import_repository, "_checkpoint", fail)
    with pytest.raises(ImportPersistenceError, match="forced failure"):
        apply(import_case, checkpoint)
    backup = import_case["root"] / "backups" / f"backup-{checkpoint}.db"
    report = import_case["root"] / "reports" / f"execution-{checkpoint}" / "import-execution-summary.json"
    assert backup.exists()
    failed = json.loads(report.read_text(encoding="utf-8"))
    assert failed["status"] == "FAILED_ROLLED_BACK" and failed["rollback_performed"] is True
    assert failed["commit_completed"] is False
    assert sha(import_case["database"]) == before
    with sqlite3.connect(import_case["database"]) as connection:
        assert connection.execute("select count(*) from work_masters").fetchone()[0] == 0
        assert connection.execute("select count(*) from work_items").fetchone()[0] == 1
        assert connection.execute("select count(*) from materials").fetchone()[0] == 6
        assert connection.execute("select count(*) from import_batches").fetchone()[0] == 0
        assert connection.execute("select count(*) from audit_logs").fetchone()[0] == 0


def test_preflight_rejects_confirmation_hash_revision_wal_and_tampered_plan_without_backup(import_case):
    base = dict(
        file_path=import_case["workbook"], project_id=PROJECT_ID,
        database_path=import_case["database"], plan_path=import_case["plan"],
        expected_workbook_sha256=sha(import_case["workbook"]),
        expected_database_sha256=sha(import_case["database"]),
        backup_path=import_case["root"] / "backups" / "never.db",
        report_dir=import_case["root"] / "reports" / "never",
        confirm=IMPORT_CONFIRMATION, repository_root=import_case["root"] / "repository",
    )
    for update, message in (
        ({"confirm": "wrong"}, "Confirmation"),
        ({"expected_workbook_sha256": "0" * 64}, "Workbook SHA256"),
        ({"expected_database_sha256": "0" * 64}, "Database SHA256"),
    ):
        with pytest.raises(ImportPreflightError, match=message):
            apply_unified_import(**{**base, **update})
        assert not base["backup_path"].exists()
    plan = json.loads(import_case["plan"].read_text(encoding="utf-8"))
    plan["work_master_summary"]["create"] = 64
    tampered = import_case["root"] / "plans" / "tampered.json"
    tampered.write_text(json.dumps(plan), encoding="utf-8")
    with pytest.raises(ImportPreflightError, match="stale or tampered"):
        apply_unified_import(**{**base, "plan_path": tampered})
    sidecar = Path(str(import_case["database"]) + "-wal")
    sidecar.touch()
    with pytest.raises(ImportPreflightError, match="WAL/SHM"):
        apply_unified_import(**base)
    sidecar.unlink()
    with sqlite3.connect(import_case["database"]) as connection:
        connection.execute("update alembic_version set version_num='0003_work_equipment_links'")
    base["expected_database_sha256"] = sha(import_case["database"])
    with pytest.raises(ImportPreflightError, match="required Alembic revision"):
        apply_unified_import(**base)


def test_cli_modes_are_mutually_exclusive_before_paths_are_opened():
    parser = build_parser()
    common = ["--file", "missing.xlsx", "--project-id", "P", "--database-path", "missing.db", "--report-dir", "reports"]
    with pytest.raises(SystemExit):
        parser.parse_args(common)
    with pytest.raises(SystemExit):
        parser.parse_args(common + ["--dry-run", "--apply"])


def test_unsafe_or_existing_backup_and_report_paths_fail_before_write(import_case):
    common = dict(
        file_path=import_case["workbook"], project_id=PROJECT_ID,
        database_path=import_case["database"], plan_path=import_case["plan"],
        expected_workbook_sha256=sha(import_case["workbook"]),
        expected_database_sha256=sha(import_case["database"]),
        confirm=IMPORT_CONFIRMATION, repository_root=import_case["root"] / "repository",
    )
    unsafe_backup = import_case["database"].parent / "backup.db"
    with pytest.raises(ImportPreflightError, match="outside"):
        apply_unified_import(
            **common, backup_path=unsafe_backup,
            report_dir=import_case["root"] / "reports" / "safe-one",
        )
    existing_backup = import_case["root"] / "backups" / "existing.db"
    existing_backup.touch()
    with pytest.raises(ImportPreflightError, match="new"):
        apply_unified_import(
            **common, backup_path=existing_backup,
            report_dir=import_case["root"] / "reports" / "safe-two",
        )
    existing_report = import_case["root"] / "reports" / "existing"
    existing_report.mkdir()
    with pytest.raises(ImportPreflightError, match="must be new"):
        apply_unified_import(
            **common, backup_path=import_case["root"] / "backups" / "safe.db",
            report_dir=existing_report,
        )
    assert not unsafe_backup.exists()


def test_conflict_plan_fails_without_backup_or_domain_write(import_case):
    with sqlite3.connect(import_case["database"]) as connection:
        connection.execute("update work_items set name='Conflict' where work_id=?", (f"{PROJECT_ID}-WRK-001",))
    conflict_plan_dir = import_case["root"] / "plans" / "conflict"
    preview = run_preview(
        file_path=import_case["workbook"], project_id=PROJECT_ID,
        database_path=import_case["database"], report_dir=conflict_plan_dir,
        repository_root=import_case["root"] / "repository",
    )
    before = sha(import_case["database"])
    backup = import_case["root"] / "backups" / "conflict.db"
    with pytest.raises(ImportPreflightError, match="contains conflicts"):
        apply_unified_import(
            file_path=import_case["workbook"], project_id=PROJECT_ID,
            database_path=import_case["database"], plan_path=Path(preview.report_paths[0]),
            expected_workbook_sha256=sha(import_case["workbook"]),
            expected_database_sha256=before, backup_path=backup,
            report_dir=import_case["root"] / "reports" / "conflict",
            confirm=IMPORT_CONFIRMATION, repository_root=import_case["root"] / "repository",
        )
    assert not backup.exists() and sha(import_case["database"]) == before
    failed = json.loads((import_case["root"] / "reports" / "conflict" / "import-execution-summary.json").read_text(encoding="utf-8"))
    assert failed["status"] == "FAILED" and failed["database_writes_performed"] is False


def test_database_change_after_backup_fails_closed_and_rolls_back(import_case, monkeypatch):
    original = import_orchestrator.execute_import_transaction

    def change_then_execute(database, *args, **kwargs):
        with sqlite3.connect(database) as connection:
            connection.execute("update projects set name='Concurrent change'")
        return original(database, *args, **kwargs)

    monkeypatch.setattr(import_orchestrator, "execute_import_transaction", change_then_execute)
    with pytest.raises(ImportPersistenceError, match="Database changed after backup"):
        apply(import_case, "stale-after-backup")
    report = import_case["root"] / "reports" / "execution-stale-after-backup" / "import-execution-summary.json"
    failed = json.loads(report.read_text(encoding="utf-8"))
    assert failed["status"] == "FAILED_ROLLED_BACK"
    assert failed["commit_completed"] is False
    assert failed["rollback_performed"] is True
    with sqlite3.connect(import_case["database"]) as connection:
        assert connection.execute("select count(*) from work_masters").fetchone()[0] == 0


def test_post_commit_verification_failure_reports_committed_state(import_case, monkeypatch):
    original = import_orchestrator._database_state
    calls = 0

    def fail_after_commit(path):
        nonlocal calls
        calls += 1
        if calls >= 3 and path == import_case["database"]:
            raise ImportPreflightError("forced verification failure")
        return original(path)

    monkeypatch.setattr(import_orchestrator, "_database_state", fail_after_commit)
    with pytest.raises(ImportPostCommitError) as raised:
        apply(import_case, "verification-failure")
    summary = raised.value.result.summary
    assert summary["status"] == "COMMITTED_VERIFICATION_FAILED"
    assert summary["commit_completed"] is True
    assert summary["database_writes_performed"] is True
    assert summary["rollback_performed"] is False
    with sqlite3.connect(import_case["database"]) as connection:
        assert connection.execute("select count(*) from work_masters").fetchone()[0] == 65


def test_post_commit_report_failure_reports_committed_state(import_case, monkeypatch):
    monkeypatch.setattr(
        import_orchestrator, "write_execution_reports",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("forced report failure")),
    )
    with pytest.raises(ImportPostCommitError) as raised:
        apply(import_case, "report-failure")
    summary = raised.value.result.summary
    assert summary["status"] == "COMMITTED_VERIFICATION_FAILED"
    assert summary["commit_completed"] is True
    assert summary["database_writes_performed"] is True
    assert summary["rollback_performed"] is False
    with sqlite3.connect(import_case["database"]) as connection:
        assert connection.execute("select count(*) from work_masters").fetchone()[0] == 65
