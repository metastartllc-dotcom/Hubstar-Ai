"""Safety and contract tests for the Work Master material recipe preview."""

import csv
import hashlib
import json
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import app.importers.work_master_material_recipe_preview as preview_module
from app.cli.preview_work_master_material_recipes import main
from app.importers.work_master_material_recipe_preview import (
    ALIAS_HEADERS,
    MATERIAL_HEADERS,
    RECIPE_HEADERS,
    WORK_HEADERS,
    RecipePreviewValidationError,
    run_recipe_preview,
)


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def _paths(tmp_path: Path) -> tuple[Path, Path, Path, Path]:
    workbook = tmp_path / "input" / "canonical.xlsx"
    database = tmp_path / "database" / "preview.db"
    repository = tmp_path / "repository"
    report = tmp_path / "reports"
    workbook.parent.mkdir()
    database.parent.mkdir()
    repository.mkdir()
    return workbook, database, repository, report


def _write_workbook(path: Path) -> None:
    book = Workbook()
    book.remove(book.active)
    work_sheet = book.create_sheet("work_masters")
    work_sheet.append(WORK_HEADERS)
    for index in range(1, 107):
        work_sheet.append([
            f"WKM-{index:06d}", f"Work {index}", "Category", "m²", 1000,
            "ACTIVE", "FIXTURE", f"WRK-{index:03d}", 100, "FIXTURE",
        ])
    material_sheet = book.create_sheet("materials")
    material_sheet.append(MATERIAL_HEADERS)
    for index in range(1, 868):
        material_sheet.append([
            f"MAT-{index:06d}", None, None, f"Material {index}", None,
            "kg", 100, "ACTIVE", "FIXTURE", 1,
        ])
    recipe_sheet = book.create_sheet("work_material_norms")
    recipe_sheet.append(RECIPE_HEADERS)
    for index in range(1, 734):
        status = "ACTIVE" if index <= 654 else "NEEDS_REVIEW"
        recipe_sheet.append([
            f"WKM-{((index - 1) % 106) + 1:06d}", f"MAT-{index:06d}",
            1 if status == "ACTIVE" or index <= 659 else None, 0, 1, 100, 100,
            status, "FIXTURE", str(index + 1), 1,
        ])
    alias_sheet = book.create_sheet("id_aliases")
    alias_sheet.append(ALIAS_HEADERS)
    for index in range(1, 166):
        alias_sheet.append([
            "MATERIAL", f"OLD-MAT-{index:06d}", f"MAT-{index:06d}", "fixture",
        ])
    book.save(path)


def _write_database(path: Path, *, full_catalog: bool = True) -> None:
    connection = sqlite3.connect(path)
    connection.executescript(
        """
        CREATE TABLE alembic_version (version_num TEXT PRIMARY KEY);
        INSERT INTO alembic_version VALUES ('0004_work_masters');
        CREATE TABLE projects (id INTEGER PRIMARY KEY, project_id TEXT NOT NULL);
        CREATE TABLE work_masters (
            id INTEGER PRIMARY KEY, work_master_id TEXT NOT NULL UNIQUE
        );
        CREATE TABLE work_items (
            id INTEGER PRIMARY KEY, project_id INTEGER NOT NULL,
            work_id TEXT NOT NULL, work_master_ref_id INTEGER
        );
        CREATE TABLE materials (
            id INTEGER PRIMARY KEY, material_id TEXT NOT NULL UNIQUE
        );
        CREATE TABLE work_material_links (
            id INTEGER PRIMARY KEY, work_id INTEGER NOT NULL, material_id INTEGER NOT NULL,
            consumption_rate REAL, waste_percentage REAL,
            calculated_quantity REAL, approved_quantity REAL, status TEXT
        );
        INSERT INTO projects VALUES (1, 'PRJ-FIXTURE');
        """
    )
    work_count = 106 if full_catalog else 1
    material_count = 867 if full_catalog else 6
    connection.executemany(
        "INSERT INTO work_masters VALUES (?, ?)",
        [(index, f"WKM-{index:06d}") for index in range(1, work_count + 1)],
    )
    connection.executemany(
        "INSERT INTO materials VALUES (?, ?)",
        [(index, f"MAT-{index:06d}") for index in range(1, material_count + 1)],
    )
    connection.execute(
        "INSERT INTO work_items VALUES (1, 1, 'PRJ-FIXTURE-WRK-001', 1)"
    )
    connection.executemany(
        "INSERT INTO work_material_links VALUES (?, 1, ?, ?, 0, ?, NULL, 'ACTIVE')",
        [(index, index, float(index), 100.0 * index) for index in range(1, 7)],
    )
    connection.commit()
    connection.close()


@pytest.fixture
def preview_inputs(tmp_path: Path) -> tuple[Path, Path, Path, Path]:
    paths = _paths(tmp_path)
    _write_workbook(paths[0])
    _write_database(paths[1])
    return paths


def test_preview_generates_deterministic_read_only_reports(preview_inputs: tuple[Path, Path, Path, Path]) -> None:
    workbook, database, repository, report = preview_inputs
    workbook_hash = _hash(workbook)
    database_hash = _hash(database)
    result = run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)

    assert result.summary["counts"] == {
        "work_masters": 106,
        "materials": 867,
        "recipes": 733,
        "source_active": 654,
        "source_needs_review": 79,
        "active_valid": 654,
        "excluded_review": 79,
        "material_aliases": 165,
        "internal_recipe_conflicts": 0,
        "invalid_aliases": 0,
        "internal_conflicts_total": 0,
        "protected_project_links": 6,
        "importable_active": 654,
    }
    assert result.summary["production_readiness"]["all_rows"] == {
        "ready": 654, "production_not_ready": 79,
    }
    assert result.summary["production_reference_readiness"]["all_rows"] == {
        "ready": 733, "production_not_ready": 0,
    }
    assert all(result.summary["reconciliation"].values())
    assert all(row.unit_conversion_applied is False for row in result.recipe_rows)
    assert len(result.report_paths) == 7
    assert _hash(workbook) == workbook_hash
    assert _hash(database) == database_hash
    for path in result.report_paths[1:]:
        assert Path(path).read_bytes().startswith(b"\xef\xbb\xbf")

    second = report.parent / "reports-2"
    run_recipe_preview(workbook, "PRJ-FIXTURE", database, second, repository)
    assert [path.read_bytes() for path in sorted(report.iterdir())] == [
        path.read_bytes() for path in sorted(second.iterdir())
    ]


def test_needs_review_is_excluded_without_inventing_values(preview_inputs: tuple[Path, Path, Path, Path]) -> None:
    workbook, database, repository, report = preview_inputs
    result = run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)
    review = [row for row in result.recipe_rows if row.source_status == "NEEDS_REVIEW"]
    assert len(review) == 79
    assert {row.action for row in review} == {"EXCLUDED_REVIEW"}
    assert sum(row.consumption_rate is None for row in review) == 74
    assert sum(row.consumption_rate is not None for row in review) == 5
    assert all(row.production_readiness == "PRODUCTION_NOT_READY" for row in review)


def test_production_not_ready_is_separate_from_internal_conflicts(tmp_path: Path) -> None:
    workbook, database, repository, report = _paths(tmp_path)
    _write_workbook(workbook)
    _write_database(database, full_catalog=False)
    result = run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)
    assert result.summary["counts"]["internal_conflicts_total"] == 0
    assert result.summary["production_reference_readiness"]["all_rows"] == {
        "ready": 1, "production_not_ready": 732,
    }
    assert result.summary["production_reference_readiness"]["active_rows"] == {
        "ready": 1, "production_not_ready": 653,
    }


@pytest.mark.parametrize(
    ("column", "value", "reason"),
    [
        (3, 0, "ACTIVE_CONSUMPTION_RATE_MUST_BE_FINITE_AND_POSITIVE"),
        (3, -1, "ACTIVE_CONSUMPTION_RATE_MUST_BE_FINITE_AND_POSITIVE"),
        (3, float("nan"), "ACTIVE_CONSUMPTION_RATE_MUST_BE_FINITE_AND_POSITIVE"),
        (3, float("inf"), "ACTIVE_CONSUMPTION_RATE_MUST_BE_FINITE_AND_POSITIVE"),
        (3, float("-inf"), "ACTIVE_CONSUMPTION_RATE_MUST_BE_FINITE_AND_POSITIVE"),
        (4, -1, "WASTE_PERCENTAGE_MUST_BE_BETWEEN_0_AND_100"),
        (4, 101, "WASTE_PERCENTAGE_MUST_BE_BETWEEN_0_AND_100"),
        (11, 0, "SOURCE_LINK_COUNT_MUST_BE_POSITIVE"),
    ],
)
def test_invalid_recipe_data_is_reported_as_internal_conflict(
    tmp_path: Path, column: int, value: object, reason: str
) -> None:
    workbook, database, repository, report = _paths(tmp_path)
    _write_workbook(workbook)
    _write_database(database)
    book = load_workbook(workbook)
    book["work_material_norms"].cell(2, column, value)
    book.save(workbook)
    book.close()
    result = run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)
    assert reason in result.recipe_rows[0].internal_conflict_reason
    assert result.recipe_rows[0].action == "INVALID"
    assert result.recipe_rows[0].production_reference_status == "READY"
    assert result.recipe_rows[0].production_readiness == "PRODUCTION_NOT_READY"
    assert result.summary["status"] == "BLOCKED_INTERNAL_CONFLICTS"


@pytest.mark.parametrize(("column", "value"), [(3, True), (3, False), (4, True), (4, False)])
def test_boolean_numeric_values_fail_closed(
    tmp_path: Path, column: int, value: bool
) -> None:
    workbook, database, repository, report = _paths(tmp_path)
    _write_workbook(workbook)
    _write_database(database)
    book = load_workbook(workbook)
    book["work_material_norms"].cell(2, column, value)
    book.save(workbook)
    book.close()
    with pytest.raises(RecipePreviewValidationError, match="Boolean"):
        run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)


@pytest.mark.parametrize("waste", [0, 100])
def test_waste_boundaries_are_valid(tmp_path: Path, waste: int) -> None:
    workbook, database, repository, report = _paths(tmp_path)
    _write_workbook(workbook)
    _write_database(database)
    book = load_workbook(workbook)
    book["work_material_norms"].cell(2, 4, waste)
    book.save(workbook)
    book.close()
    result = run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)
    assert result.recipe_rows[0].internal_conflict_reason == ""


def test_duplicate_pair_and_invalid_alias_fail_closed_in_conflict_report(preview_inputs: tuple[Path, Path, Path, Path]) -> None:
    workbook, database, repository, report = preview_inputs
    book = load_workbook(workbook)
    recipe = book["work_material_norms"]
    recipe.cell(3, 1, recipe.cell(2, 1).value)
    recipe.cell(3, 2, recipe.cell(2, 2).value)
    alias = book["id_aliases"]
    alias.cell(2, 3, "MISSING")
    book.save(workbook)
    book.close()
    result = run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)
    assert result.summary["counts"]["internal_recipe_conflicts"] == 2
    assert result.summary["counts"]["invalid_aliases"] == 1
    with (report / "work-master-material-recipe-conflicts.csv").open(
        encoding="utf-8-sig", newline=""
    ) as source:
        conflicts = list(csv.DictReader(source))
    assert len(conflicts) == 3


@pytest.mark.parametrize(
    ("alias_id", "canonical_id", "reason"),
    [
        ("MAT-000867", "MAT-000001", "ALIAS_COLLIDES_WITH_CANONICAL_MATERIAL_ID"),
        ("OLD-MAT-000001", "OLD-MAT-000002", "ALIAS_CHAIN_OR_CYCLE_NOT_ALLOWED"),
        ("OLD-MAT-000001", "MISSING", "MISSING_CANONICAL_MATERIAL"),
    ],
)
def test_alias_collision_cycle_and_missing_target_are_invalid(
    tmp_path: Path, alias_id: str, canonical_id: str, reason: str
) -> None:
    workbook, database, repository, report = _paths(tmp_path)
    _write_workbook(workbook)
    _write_database(database)
    book = load_workbook(workbook)
    alias = book["id_aliases"]
    alias.cell(2, 2, alias_id)
    alias.cell(2, 3, canonical_id)
    book.save(workbook)
    book.close()
    result = run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)
    assert reason in result.alias_rows[0].conflict_reason
    assert result.alias_rows[0].validation_status == "INVALID"


def test_alias_cycle_is_rejected(tmp_path: Path) -> None:
    workbook, database, repository, report = _paths(tmp_path)
    _write_workbook(workbook)
    _write_database(database)
    book = load_workbook(workbook)
    alias = book["id_aliases"]
    first = alias.cell(2, 2).value
    second = alias.cell(3, 2).value
    alias.cell(2, 3, second)
    alias.cell(3, 3, first)
    book.save(workbook)
    book.close()
    result = run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)
    invalid = [row for row in result.alias_rows if row.validation_status == "INVALID"]
    assert len(invalid) == 2
    assert all("ALIAS_CHAIN_OR_CYCLE_NOT_ALLOWED" in row.conflict_reason for row in invalid)


def test_report_location_and_overwrite_are_rejected(preview_inputs: tuple[Path, Path, Path, Path]) -> None:
    workbook, database, repository, report = preview_inputs
    with pytest.raises(RecipePreviewValidationError, match="outside repository"):
        run_recipe_preview(workbook, "PRJ-FIXTURE", database, repository / "reports", repository)
    report.mkdir()
    (report / "unrelated.txt").write_text("occupied", encoding="utf-8")
    with pytest.raises(RecipePreviewValidationError, match="new or empty"):
        run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)


def test_cli_requires_dry_run_before_opening_inputs(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as exc:
        main([
            "--file", "missing.xlsx", "--project-id", "PRJ",
            "--database-path", "missing.db", "--report-dir", "missing-reports",
        ])
    assert exc.value.code == 2
    assert "--dry-run is required" in capsys.readouterr().err


def test_cli_writes_reports_then_returns_nonzero_for_blocking_conflict(
    preview_inputs: tuple[Path, Path, Path, Path],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook, database, repository, report = preview_inputs
    book = load_workbook(workbook)
    book["work_material_norms"].cell(2, 3, 0)
    book.save(workbook)
    book.close()
    monkeypatch.chdir(repository)
    exit_code = main([
        "--file", str(workbook), "--project-id", "PRJ-FIXTURE",
        "--database-path", str(database), "--report-dir", str(report), "--dry-run",
    ])
    assert exit_code == 1
    assert len(list(report.iterdir())) == 7


@pytest.mark.parametrize(
    ("changed_call", "message"),
    [(3, "workbook changed"), (4, "Database changed")],
)
def test_hash_change_is_detected_after_reports(
    preview_inputs: tuple[Path, Path, Path, Path],
    monkeypatch: pytest.MonkeyPatch,
    changed_call: int,
    message: str,
) -> None:
    workbook, database, repository, report = preview_inputs
    actual = preview_module.file_sha256
    calls = 0

    def changed_after_read(path: Path) -> str:
        nonlocal calls
        calls += 1
        value = actual(path)
        return "0" * 64 if calls == changed_call else value

    monkeypatch.setattr(preview_module, "file_sha256", changed_after_read)
    with pytest.raises(RecipePreviewValidationError, match=message):
        run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)


def test_read_only_connection_is_closed(preview_inputs: tuple[Path, Path, Path, Path]) -> None:
    workbook, database, repository, report = preview_inputs
    run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)
    database.unlink()
    assert not database.exists()


def test_summary_json_and_rows_are_stably_sorted(preview_inputs: tuple[Path, Path, Path, Path]) -> None:
    workbook, database, repository, report = preview_inputs
    result = run_recipe_preview(workbook, "PRJ-FIXTURE", database, report, repository)
    parsed = json.loads((report / "work-master-material-recipe-summary.json").read_text("utf-8"))
    assert parsed == result.summary
    pairs = [(row.work_master_id, row.material_id) for row in result.recipe_rows]
    assert pairs == sorted(pairs)
