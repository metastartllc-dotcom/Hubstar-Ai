import csv
import hashlib
import json
import sqlite3
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.cli import import_excel
from app.importers.unified_excel_preview import (
    MATERIAL_HEADERS,
    PreviewValidationError,
    WORK_HEADERS,
    _material_preview,
    _number,
    run_preview,
)
from app.repositories import import_preview as preview_repository


PROJECT_ID = "PRJ-ALTAI-R7-B"


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _material_id(index: int) -> str:
    prefix = "PRD-ALTAI-B" if index <= 138 else "PRD-ALTAI-6R"
    return f"{prefix}-{index:04d}"


def make_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ("Заавар", "Материалын үнэ", "Ажил-Тээвэр-Механизм", "Нэгтгэлийн хяналт"):
        workbook.create_sheet(name)
    material_sheet = workbook["Материалын үнэ"]
    work_sheet = workbook["Ажил-Тээвэр-Механизм"]
    for column, value in enumerate(MATERIAL_HEADERS, start=1):
        material_sheet.cell(4, column, value)
    for column, value in enumerate(WORK_HEADERS, start=1):
        work_sheet.cell(4, column, value)

    for index in range(1, 987):
        name = f"Материал {index}"
        price = float(index * 10) if index <= 810 else None
        if index == 811:
            name = "Нэг үнэ бүхий ижил материал"
        elif index == 1:
            name = "Нэг үнэ бүхий ижил материал"
        elif index in (2, 3, 812):
            name = "Олон үнэ бүхий ижил материал"
            price = {2: 100.0, 3: 300.0, 812: None}[index]
        row = [
            _material_id(index), "Тест ангилал", name, "м²", None,
            price, None, None, None, None,
        ]
        material_sheet.append(row)

    for index in range(1, 66):
        work_sheet.append([
            f"WRK-ALTAI-B-{index:03d}", "Ангилал", f"Ажил {index}",
            float(index), "м²", 1000.0, None, None, None, None, None,
            None, None, None, None,
        ])
    workbook.save(path)


def make_database(path: Path, *, work_conflict: bool = False, material_conflict: bool = False) -> None:
    connection = sqlite3.connect(path)
    connection.executescript(
        """
        CREATE TABLE projects (id INTEGER PRIMARY KEY, project_id TEXT UNIQUE, name TEXT NOT NULL);
        CREATE TABLE work_items (
          id INTEGER PRIMARY KEY, work_id TEXT UNIQUE, project_id INTEGER,
          wbs_code TEXT, name TEXT NOT NULL, unit TEXT, quantity REAL,
          labor_unit_rate REAL, labor_total REAL, status TEXT
        );
        CREATE TABLE materials (
          id INTEGER PRIMARY KEY, material_id TEXT UNIQUE, master_id TEXT,
          code TEXT, name TEXT NOT NULL, specification TEXT,
          normalized_unit TEXT, supplier_id INTEGER, unit_price REAL, status TEXT
        );
        """
    )
    connection.execute("INSERT INTO projects VALUES (1, ?, ?)", (PROJECT_ID, "Тест төсөл"))
    connection.execute(
        "INSERT INTO work_items VALUES (1, ?, 1, NULL, ?, 'м²', 1, 1000, 1000, 'ACTIVE')",
        (f"{PROJECT_ID}-WRK-001", "Зөрүү" if work_conflict else "Ажил 1"),
    )
    for offset, index in enumerate(range(981, 987), start=1):
        connection.execute(
            "INSERT INTO materials VALUES (?, ?, NULL, NULL, ?, NULL, 'м²', NULL, ?, 'ACTIVE')",
            (
                offset, _material_id(index),
                "Зөрүү" if material_conflict and index == 981 else f"Материал {index}",
                float(index),
            ),
        )
    connection.commit()
    connection.close()


@pytest.fixture
def preview_inputs(tmp_path: Path):
    workbook = tmp_path / "synthetic.xlsx"
    database = tmp_path / "synthetic.db"
    make_workbook(workbook)
    make_database(database)
    reports = tmp_path.parent / f"reports-{tmp_path.name}"
    return workbook, database, reports


def test_complete_preview_is_read_only_deterministic_and_utf8(preview_inputs, tmp_path):
    workbook, database, reports = preview_inputs
    workbook_before, database_before = _hash(workbook), _hash(database)
    result = run_preview(
        file_path=workbook, project_id=PROJECT_ID, database_path=database,
        report_dir=reports, repository_root=tmp_path / "repository",
        generated_at_utc="2026-01-01T00:00:00Z",
    )

    assert result.summary["work_summary"] == {
        "source": 65, "valid": 65, "create": 64,
        "existing_identical": 1, "existing_conflict": 0,
    }
    materials = result.summary["material_summary"]
    assert materials["source"] == materials["valid"] == 986
    assert materials["create"] == 980
    assert materials["existing_identical"] == 6
    assert materials["existing_conflict"] == 0
    assert materials["excel_price"] == 810
    assert materials["db_price_preserved"] == 6
    assert result.summary["database_writes_performed"] is False
    assert result.summary["database"]["sqlite_total_changes"] == 0
    assert _hash(workbook) == workbook_before
    assert _hash(database) == database_before
    assert [row.canonical_work_id for row in result.work_rows] == sorted(
        row.canonical_work_id for row in result.work_rows
    )
    assert [row.material_id for row in result.material_rows] == sorted(
        row.material_id for row in result.material_rows
    )
    assert len(result.report_paths) == 5
    assert Path(result.report_paths[1]).read_bytes().startswith(b"\xef\xbb\xbf")
    assert "Материал" in Path(result.report_paths[2]).read_text(encoding="utf-8-sig")
    summary = json.loads(Path(result.report_paths[0]).read_text(encoding="utf-8"))
    assert summary["generated_at_utc"] == "2026-01-01T00:00:00Z"


def test_price_precedence_unique_ambiguous_unresolved_and_db_preservation(preview_inputs, tmp_path):
    workbook, database, reports = preview_inputs
    result = run_preview(
        file_path=workbook, project_id=PROJECT_ID, database_path=database,
        report_dir=reports, repository_root=tmp_path / "repository",
    )
    rows = {row.material_id: row for row in result.material_rows}
    assert rows[_material_id(1)].price_resolution == "excel_current"
    assert rows[_material_id(811)].price_resolution == "exact_reference_unique"
    assert rows[_material_id(811)].proposed_unit_price == 10
    assert rows[_material_id(812)].price_resolution == "exact_reference_ambiguous"
    assert rows[_material_id(812)].candidate_price == 200
    assert rows[_material_id(812)].proposed_unit_price is None
    assert rows[_material_id(813)].price_resolution == "unresolved"
    assert rows[_material_id(813)].proposed_status == "NEEDS_REVIEW"
    assert rows[_material_id(981)].price_resolution == "db_price_preserved"
    assert rows[_material_id(981)].proposed_unit_price == 981
    assert rows[_material_id(981)].action == "SKIP"
    assert rows[_material_id(811)].category == "Тест ангилал"
    assert rows[_material_id(811)].specification is None


def test_new_price_precedes_current_and_reference_requires_compatible_unit():
    parsed = [
        {
            "material_id": "M-1", "name": "Ижил", "normalized_unit": "м²",
            "specification": None, "category": None, "current_price": 100.0,
            "new_price": 120.0, "source_row": 5,
        },
        {
            "material_id": "M-2", "name": "Ижил", "normalized_unit": "кг",
            "specification": None, "category": None, "current_price": None,
            "new_price": None, "source_row": 6,
        },
    ]
    rows = {row.material_id: row for row in _material_preview(
        parsed, SimpleNamespace(materials={})
    )}
    assert rows["M-1"].source_price_kind == "excel_new"
    assert rows["M-1"].proposed_unit_price == 120
    assert rows["M-2"].price_resolution == "unresolved"
    assert rows["M-2"].proposed_unit_price is None


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        ("hide_sheet", "not visible"),
        ("bad_header", "header"),
        ("remove_material", "Expected 986"),
        ("duplicate_material", "Duplicate Product"),
        ("duplicate_work", "Duplicate Work"),
        ("bad_work_id", "Invalid Work ID"),
    ],
)
def test_template_validation_fails_before_reports(preview_inputs, tmp_path, mutation, message):
    workbook_path, database, reports = preview_inputs
    workbook = load_workbook(workbook_path)
    if mutation == "hide_sheet":
        workbook["Заавар"].sheet_state = "hidden"
    elif mutation == "bad_header":
        workbook["Материалын үнэ"].cell(4, 1, "Wrong")
    elif mutation == "remove_material":
        workbook["Материалын үнэ"].cell(990, 1).value = None
    elif mutation == "duplicate_material":
        workbook["Материалын үнэ"].cell(990, 1, _material_id(1))
    elif mutation == "duplicate_work":
        workbook["Ажил-Тээвэр-Механизм"].cell(69, 1, "WRK-ALTAI-B-001")
    else:
        workbook["Ажил-Тээвэр-Механизм"].cell(5, 1, "BAD")
    workbook.save(workbook_path)
    with pytest.raises(PreviewValidationError, match=message):
        run_preview(
            file_path=workbook_path, project_id=PROJECT_ID, database_path=database,
            report_dir=reports, repository_root=tmp_path / "repository",
        )
    assert not reports.exists()


def test_existing_conflicts_are_reported(preview_inputs, tmp_path):
    workbook, database, reports = preview_inputs
    database.unlink()
    make_database(database, work_conflict=True, material_conflict=True)
    result = run_preview(
        file_path=workbook, project_id=PROJECT_ID, database_path=database,
        report_dir=reports, repository_root=tmp_path / "repository",
    )
    assert result.summary["work_summary"]["existing_conflict"] == 1
    assert result.summary["material_summary"]["existing_conflict"] == 1
    with Path(result.report_paths[4]).open(encoding="utf-8-sig", newline="") as source:
        conflicts = list(csv.DictReader(source))
    assert [(row["entity_type"], row["external_id"]) for row in conflicts] == sorted(
        (row["entity_type"], row["external_id"]) for row in conflicts
    )


@pytest.mark.parametrize("value", [True, -1, float("nan"), float("inf"), float("-inf")])
def test_invalid_numeric_values_are_rejected(value):
    with pytest.raises(PreviewValidationError):
        _number(value)


def test_blank_numeric_is_none_not_nan_string():
    assert _number(None) is None
    assert _number("   ") is None


def test_unsafe_or_nonempty_report_directory_is_rejected(preview_inputs, tmp_path):
    workbook, database, _ = preview_inputs
    unsafe = tmp_path / "inside-repository"
    with pytest.raises(PreviewValidationError, match="outside"):
        run_preview(
            file_path=workbook, project_id=PROJECT_ID, database_path=database,
            report_dir=unsafe, repository_root=tmp_path,
        )
    existing = tmp_path.parent / f"existing-{tmp_path.name}"
    existing.mkdir()
    (existing / "keep.txt").write_text("keep", encoding="utf-8")
    with pytest.raises(PreviewValidationError, match="empty"):
        run_preview(
            file_path=workbook, project_id=PROJECT_ID, database_path=database,
            report_dir=existing, repository_root=tmp_path,
        )


def test_cli_without_dry_run_fails_before_preview(monkeypatch):
    called = False

    def forbidden(**_kwargs):
        nonlocal called
        called = True

    monkeypatch.setattr(import_excel, "run_preview", forbidden)
    with pytest.raises(SystemExit) as exc:
        import_excel.main([
            "--file", "missing.xlsx", "--project-id", PROJECT_ID,
            "--database-path", "missing.db", "--report-dir", "reports",
        ])
    assert exc.value.code == 2
    assert called is False


def test_database_is_opened_with_uri_mode_ro(monkeypatch, preview_inputs):
    _, database, _ = preview_inputs
    real_connect = sqlite3.connect
    calls = []

    def recording_connect(database_arg, *args, **kwargs):
        calls.append((database_arg, kwargs.copy()))
        return real_connect(database_arg, *args, **kwargs)

    monkeypatch.setattr(preview_repository.sqlite3, "connect", recording_connect)
    snapshot = preview_repository.load_database_snapshot(database, PROJECT_ID)
    assert snapshot.total_changes == 0
    assert calls and "mode=ro" in calls[0][0]
    assert calls[0][1]["uri"] is True


def test_work_mapping_uses_new_rate_and_never_leaks_category(preview_inputs, tmp_path):
    workbook_path, database, reports = preview_inputs
    workbook = load_workbook(workbook_path)
    sheet = workbook["Ажил-Тээвэр-Механизм"]
    sheet.cell(6, 7, 2222)
    workbook.save(workbook_path)
    result = run_preview(
        file_path=workbook_path, project_id=PROJECT_ID, database_path=database,
        report_dir=reports, repository_root=tmp_path / "repository",
    )
    row = next(item for item in result.work_rows if item.source_work_id.endswith("002"))
    assert row.canonical_work_id == f"{PROJECT_ID}-WRK-002"
    assert row.labor_unit_rate == 2222
    assert row.category == "Ангилал"
    assert "wbs_code" not in row.to_dict()
