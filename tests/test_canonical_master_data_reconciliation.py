"""Tests for canonical workbook/production reconciliation preview."""

import hashlib
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.cli.reconcile_canonical_master_data import main
from app.importers.canonical_master_data_reconciliation import run_reconciliation_preview
from app.importers.work_master_material_recipe_preview import (
    ALIAS_HEADERS, MATERIAL_HEADERS, RECIPE_HEADERS, WORK_HEADERS,
)


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _paths(tmp_path: Path) -> tuple[Path, Path, Path, Path]:
    wb, db = tmp_path / "input" / "canonical.xlsx", tmp_path / "database" / "test.db"
    repo, report = tmp_path / "repository", tmp_path / "reports"
    wb.parent.mkdir(); db.parent.mkdir(); repo.mkdir()
    return wb, db, repo, report


def _workbook(path: Path) -> None:
    book = Workbook(); book.remove(book.active)
    works = book.create_sheet("work_masters"); works.append(WORK_HEADERS)
    for i in range(1, 107):
        works.append([f"WKM-{i:06d}", f"Work {i}", "Category", "m²", 10, "ACTIVE", "FIXTURE", f"WRK-{i:03d}", 100, "fixture"])
    materials = book.create_sheet("materials"); materials.append(MATERIAL_HEADERS)
    for i in range(1, 868):
        materials.append([f"MAT-{i:06d}", f"MASTER-{i}", f"CODE-{i}", f"Material {i}", f"Spec {i}", "kg", 100, "ACTIVE", "fixture", 1])
    recipes = book.create_sheet("work_material_norms"); recipes.append(RECIPE_HEADERS)
    for i in range(1, 734):
        status = "ACTIVE" if i <= 654 else "NEEDS_REVIEW"
        recipes.append([f"WKM-{((i-1)%106)+1:06d}", f"MAT-{i:06d}", 1 if status == "ACTIVE" else None, 0, 1, 100, 100, status, "FIXTURE", str(i), 1])
    aliases = book.create_sheet("id_aliases"); aliases.append(ALIAS_HEADERS)
    for i in range(1, 166): aliases.append(["MATERIAL", f"OLD-{i:06d}", f"MAT-{i:06d}", "legacy"])
    for i in range(1, 45): aliases.append(["WORK_PACKAGE", f"PKG-{i:03d}", f"WKM-{i:06d}", "package mapping"])
    reviewed = book.create_sheet("materials_reviewed"); reviewed.append(["material_id", "proposed_name", "proposed_unit"]); reviewed.append(["MAT-000001", "DO NOT USE", "m³"])
    book.save(path)


def _database(path: Path) -> None:
    c = sqlite3.connect(path)
    c.executescript("""
      CREATE TABLE alembic_version(version_num TEXT PRIMARY KEY);
      INSERT INTO alembic_version VALUES('0004_work_masters');
      CREATE TABLE projects(id INTEGER PRIMARY KEY,project_id TEXT);
      CREATE TABLE work_masters(id INTEGER PRIMARY KEY,work_master_id TEXT,name TEXT,category TEXT,default_unit TEXT,default_labor_unit_rate REAL,status TEXT,source_dataset TEXT,source_work_id TEXT);
      CREATE TABLE work_items(id INTEGER PRIMARY KEY,work_id TEXT,project_id INTEGER,work_master_ref_id INTEGER,wbs_code TEXT,name TEXT,unit TEXT,quantity REAL,labor_unit_rate REAL,labor_total REAL,status TEXT);
      CREATE TABLE materials(id INTEGER PRIMARY KEY,material_id TEXT,master_id TEXT,code TEXT,name TEXT,specification TEXT,normalized_unit TEXT,unit_price REAL,status TEXT);
      CREATE TABLE work_material_links(id INTEGER PRIMARY KEY,work_id INTEGER,material_id INTEGER,consumption_rate REAL,waste_percentage REAL,calculated_quantity REAL,approved_quantity REAL,status TEXT);
      CREATE TABLE equipments(id INTEGER PRIMARY KEY,equipment_id TEXT,master_id TEXT,type TEXT,model TEXT,capacity TEXT,location TEXT,operator_included INTEGER,fuel_included INTEGER,delivery_included INTEGER,included_delivery_one_way_distance_km REAL,tariff_type TEXT,unit_rate REAL,availability TEXT,status TEXT);
      CREATE TABLE work_equipment_links(id INTEGER PRIMARY KEY,work_item_id INTEGER,equipment_id INTEGER,usage_quantity REAL,agreed_unit_rate REAL,tariff_type_snapshot TEXT,operator_included_snapshot INTEGER,fuel_included_snapshot INTEGER,delivery_included_snapshot INTEGER,included_delivery_one_way_distance_km_snapshot REAL,status TEXT);
      INSERT INTO projects VALUES(1,'PRJ-FIXTURE');
    """)
    c.executemany("INSERT INTO work_masters VALUES(?,?,?,?,?,?,?,?,?)", [(i,f"WKM-{i:06d}",f"Work {i}","Category","m²",10,"ACTIVE","FIXTURE",f"WRK-{i:03d}") for i in range(1,66)])
    canonical = [(i,f"MAT-{i:06d}",f"MASTER-{i}",f"CODE-{i}",f"Material {i}",f"Spec {i}","kg",100,"ACTIVE") for i in range(1,822)]
    legacy = [(821+i,f"OLD-{i:06d}",None,None,f"Legacy {i}",None,"kg",50,"ACTIVE") for i in range(1,166)]
    c.executemany("INSERT INTO materials VALUES(?,?,?,?,?,?,?,?,?)", canonical + legacy)
    c.execute("INSERT INTO work_items VALUES(1,'PRJ-FIXTURE-WRK-001',1,1,NULL,'Facade','m²',100,10,1000,'ACTIVE')")
    c.executemany("INSERT INTO work_material_links VALUES(?,1,?,1,0,100,NULL,'ACTIVE')", [(i,i) for i in range(1,7)])
    c.execute("INSERT INTO equipments VALUES(1,'EQP-1',NULL,'Crane',NULL,'25 t',NULL,1,1,1,25,'MNT/hour',150000,NULL,'ACTIVE')")
    c.execute("INSERT INTO work_equipment_links VALUES(1,1,1,72,150000,'MNT/hour',1,1,1,25,'ACTIVE')")
    c.commit(); c.close()


@pytest.fixture
def inputs(tmp_path: Path) -> tuple[Path, Path, Path, Path]:
    paths = _paths(tmp_path); _workbook(paths[0]); _database(paths[1]); return paths


def test_baseline_actions_protected_data_and_determinism(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb, db, repo, report = inputs; before = (_sha(wb), _sha(db))
    result = run_reconciliation_preview(wb, "PRJ-FIXTURE", db, report, repo)
    assert result.summary["actions"] == {
        "ALIAS_REVIEW_REQUIRED": 209, "CREATE_PROPOSAL": 87,
        "KEEP_IDENTICAL": 886, "RETAIN_PRODUCTION_ONLY": 165,
    }
    assert len(result.protected_rows) == 9
    assert {row.entity_type for row in result.protected_rows} == {"FACADE_WORK_ITEM", "FACADE_MATERIAL_LINK", "EQUIPMENT_MASTER", "WORK_EQUIPMENT_SNAPSHOT"}
    assert before == (_sha(wb), _sha(db))
    second = report.parent / "reports-2"
    run_reconciliation_preview(wb, "PRJ-FIXTURE", db, second, repo)
    assert [p.read_bytes() for p in sorted(report.iterdir())] == [p.read_bytes() for p in sorted(second.iterdir())]
    assert all(p.read_bytes().startswith(b"\xef\xbb\xbf") for p in report.iterdir() if p.suffix == ".csv")
    assert result.summary["counts"]["workbook_work_masters"] == 106
    assert result.summary["counts"]["workbook_materials"] == 867
    assert result.summary["counts"]["production_materials"] == 986


def test_source_collision_is_conflict(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb, db, repo, report = inputs
    connection=sqlite3.connect(db); connection.execute("UPDATE work_masters SET source_dataset='OTHER',source_work_id='OTHER-001' WHERE work_master_id='WKM-000001'"); connection.commit(); connection.close()
    book=load_workbook(wb); ws=book["work_masters"]
    ws.cell(67,7,"OTHER"); ws.cell(67,8,"OTHER-001"); book.save(wb); book.close()
    result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    row=next(row for row in result.work_rows if row.external_id=="WKM-000066")
    assert row.action=="CONFLICT" and "source_identity_owner" in row.explicit_diffs


def test_workbook_source_identity_collision_marks_both_rows_conflict(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb,db,repo,report=inputs; book=load_workbook(wb); ws=book["work_masters"]
    ws.cell(67,7,"FIXTURE"); ws.cell(67,8,"WRK-002"); book.save(wb); book.close()
    result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    rows=[row for row in result.work_rows if row.external_id in {"WKM-000002","WKM-000066"}]
    assert len(rows)==2 and all(row.action=="CONFLICT" for row in rows)
    assert all("workbook_source_identity_count" in row.explicit_diffs for row in rows)


def test_specification_ambiguity_requires_review_not_merge(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb,db,repo,report=inputs; c=sqlite3.connect(db); c.execute("UPDATE materials SET specification='Different' WHERE material_id='MAT-000001'"); c.commit(); c.close()
    result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    row=next(row for row in result.material_rows if row.external_id=="MAT-000001")
    assert row.action=="UPDATE_REVIEW_REQUIRED" and "specification" in row.explicit_diffs
    assert next(r for r in result.recipe_rows if r.material_id=="MAT-000001").readiness=="DEPENDS_ON_UPDATE_REVIEW"


def test_blank_workbook_price_preserves_existing_price(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb,db,repo,report=inputs; book=load_workbook(wb); book["materials"].cell(2,7).value=None; book.save(wb); book.close()
    result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    row=next(row for row in result.material_rows if row.external_id=="MAT-000001")
    assert row.action=="UPDATE_REVIEW_REQUIRED"
    assert "PRESERVE_PRODUCTION_NON_NULL_PRICE" in row.explicit_diffs
    protected=next(row for row in result.protected_rows if row.entity_type=="FACADE_MATERIAL_LINK" and row.external_id=="MAT-000001")
    assert "PRESERVE_PRODUCTION_NON_NULL_PRICE" in protected.snapshot_json
    assert result.summary["update_review_breakdown"]["field_difference_occurrences"]["price"]==1


@pytest.mark.parametrize(("alias_id","target","issue"), [("MAT-000867","MAT-000001","ALIAS_COLLIDES_WITH_CANONICAL_ID"),("OLD-000001","MISSING","MISSING_CANONICAL_TARGET"),("OLD-000001","OLD-000002","ALIAS_CHAIN_OR_CYCLE")])
def test_alias_target_cycle_collision(inputs: tuple[Path, Path, Path, Path], alias_id: str, target: str, issue: str) -> None:
    wb,db,repo,report=inputs; book=load_workbook(wb); ws=book["id_aliases"]; ws.cell(2,2,alias_id); ws.cell(2,3,target); book.save(wb); book.close()
    result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    row=next(row for row in result.alias_rows if row.alias_id==alias_id)
    assert row.action=="CONFLICT" and issue in row.validation_issues


def test_two_way_alias_cycle_is_conflict(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb,db,repo,report=inputs; book=load_workbook(wb); ws=book["id_aliases"]
    first,second=ws.cell(2,2).value,ws.cell(3,2).value
    ws.cell(2,3,second); ws.cell(3,3,first); book.save(wb); book.close()
    result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    rows=[row for row in result.alias_rows if row.alias_id in {first,second}]
    assert len(rows)==2 and all("ALIAS_CHAIN_OR_CYCLE" in row.validation_issues for row in rows)


def test_proposed_review_sheet_is_ignored_and_production_only_retained(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb,db,repo,report=inputs; result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    material=next(row for row in result.material_rows if row.external_id=="MAT-000001")
    assert material.action=="KEEP_IDENTICAL" and "DO NOT USE" not in material.explicit_diffs
    assert next(row for row in result.material_rows if row.external_id=="OLD-000165").action=="RETAIN_PRODUCTION_ONLY"
    package=next(row for row in result.alias_rows if row.entity_type=="WORK_PACKAGE")
    assert package.interpretation=="WORK_PACKAGE_MAPPING_NOT_WORK_MASTER_ALIAS"


def test_recipe_unit_mismatch_and_invalid_review_are_not_importable(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb,db,repo,report=inputs; c=sqlite3.connect(db); c.execute("UPDATE materials SET normalized_unit='m' WHERE material_id='MAT-000001'"); c.commit(); c.close()
    book=load_workbook(wb); book["materials"].cell(3,6).value=None; book.save(wb); book.close()
    result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    first=next(row for row in result.recipe_rows if row.material_id=="MAT-000001")
    second=next(row for row in result.recipe_rows if row.material_id=="MAT-000002")
    review=next(row for row in result.recipe_rows if row.source_status=="NEEDS_REVIEW")
    assert first.unit_match=="MISMATCH" and first.readiness=="DEPENDS_ON_UPDATE_REVIEW"
    assert second.internal_validation=="INVALID" and not second.importable
    assert review.readiness=="NOT_IMPORTABLE_REVIEW" and not review.importable


def test_missing_production_reference_is_not_semantic_match(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb,db,repo,report=inputs; result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    existing=next(row for row in result.recipe_rows if row.work_master_id=="WKM-000001" and row.material_id=="MAT-000001")
    missing=next(row for row in result.recipe_rows if not row.work_production_reference)
    assert existing.semantic_match=="MATCH"
    assert missing.semantic_match=="NOT_EVALUATED"


def test_multiple_dependencies_are_preserved_when_create_is_primary(inputs: tuple[Path, Path, Path, Path]) -> None:
    wb,db,repo,report=inputs
    connection=sqlite3.connect(db); connection.execute("UPDATE materials SET normalized_unit='m' WHERE material_id='MAT-000001'"); connection.commit(); connection.close()
    book=load_workbook(wb); book["work_material_norms"].cell(2,1,"WKM-000066"); book.save(wb); book.close()
    result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    row=next(row for row in result.recipe_rows if row.work_master_id=="WKM-000066" and row.material_id=="MAT-000001")
    assert row.readiness=="DEPENDS_ON_CREATE"
    assert set(row.dependency_actions.split(";")) >= {"CREATE_PROPOSAL","UPDATE_REVIEW_REQUIRED","ALIAS_REVIEW_REQUIRED","UNIT_REVIEW_REQUIRED"}


@pytest.mark.parametrize(("column","value"), [(3,"not-a-number"),(3,float("nan")),(3,float("inf")),(4,"bad"),(4,-1),(4,101)])
def test_malformed_recipe_numbers_are_reported_not_crashed(inputs: tuple[Path, Path, Path, Path], column: int, value: object) -> None:
    wb,db,repo,report=inputs; book=load_workbook(wb); book["work_material_norms"].cell(2,column,value); book.save(wb); book.close()
    result=run_reconciliation_preview(wb,"PRJ-FIXTURE",db,report,repo)
    assert result.recipe_rows[0].internal_validation=="INVALID"
    assert not result.recipe_rows[0].importable


def test_cli_fail_closed_and_connection_cleanup(inputs: tuple[Path, Path, Path, Path], monkeypatch: pytest.MonkeyPatch) -> None:
    wb,db,repo,report=inputs
    with pytest.raises(SystemExit) as exc: main(["--file",str(wb),"--project-id","PRJ-FIXTURE","--database-path",str(db),"--report-dir",str(report)])
    assert exc.value.code==2
    book=load_workbook(wb); book["materials"].cell(3,6).value=None; book.save(wb); book.close(); monkeypatch.chdir(repo)
    assert main(["--file",str(wb),"--project-id","PRJ-FIXTURE","--database-path",str(db),"--report-dir",str(report),"--dry-run"])==1
    assert len(list(report.iterdir()))==7
    db.unlink(); assert not db.exists()
