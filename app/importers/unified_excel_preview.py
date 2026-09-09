"""Read-only preview for the unified 6R/7R Excel import workbook."""

from __future__ import annotations

import hashlib
import math
import re
import statistics
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from pydantic import ValidationError

from app.exporters.import_preview_reports import REPORT_FILENAMES, write_preview_reports
from app.repositories.import_preview import load_database_snapshot
from app.schemas.import_preview import MaterialPreviewRow, PreviewResult, WorkPreviewRow
from app.schemas.schemas import MaterialCreateRequest, ProjectWorkItemCreate
from app.validators.units import UnitNormalizer


REQUIRED_SHEETS = (
    "Заавар",
    "Материалын үнэ",
    "Ажил-Тээвэр-Механизм",
    "Нэгтгэлийн хяналт",
)
MATERIAL_HEADERS = (
    "Product ID", "Ангилал", "Материалын нэр", "Нэгж", "Үзүүлэлт",
    "Одоогийн үнэ MNT", "ШИНЭ НЭГЖ ҮНЭ MNT", "Үнийн огноо",
    "Нийлүүлэгч/эх сурвалж", "Тайлбар",
)
WORK_HEADERS = (
    "Work ID", "Ангилал", "Ажлын нэр", "Хэмжээ", "Нэгж",
    "Суурь ажлын үнэ", "ШИНЭ АЖЛЫН ҮНЭ", "Тээврийн нийт",
    "Механизмын нийт", "НӨАТ %", "Эх сурвалж", "Ажлын хөлсний нийт",
    "НӨАТ-ын өмнөх дүн", "НӨАТ", "Нийт",
)
WORK_ID_PATTERN = re.compile(r"^WRK-ALTAI-B-(\d{3})$")
EXPECTED_MATERIAL_ROWS = 986
EXPECTED_WORK_ROWS = 65


class PreviewValidationError(ValueError):
    """Raised before reports are created when preview inputs are unsafe or invalid."""


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def validate_report_directory(
    report_dir: Path,
    repository_root: Path,
    workbook_path: Path,
    database_path: Path,
) -> Path:
    resolved = report_dir.resolve()
    forbidden = (
        repository_root.resolve(), workbook_path.resolve().parent,
        database_path.resolve().parent,
    )
    if any(_is_relative_to(resolved, root) for root in forbidden):
        raise PreviewValidationError(
            "Report directory must be outside the repository, workbook, and database directories"
        )
    if resolved.exists():
        if not resolved.is_dir():
            raise PreviewValidationError("Report path is not a directory")
        if any(resolved.iterdir()):
            raise PreviewValidationError("Report directory must be empty")
    if any((resolved / name).exists() for name in REPORT_FILENAMES):
        raise PreviewValidationError("Preview report already exists")
    return resolved


def _text(value: Any, *, required: bool = False) -> str | None:
    if value is None:
        if required:
            raise PreviewValidationError("Required text value is blank")
        return None
    if isinstance(value, bool):
        raise PreviewValidationError("Boolean is not a valid text value")
    result = unicodedata.normalize("NFKC", str(value)).strip()
    result = " ".join(result.split())
    if not result:
        if required:
            raise PreviewValidationError("Required text value is blank")
        return None
    return result


def _number(value: Any, *, required: bool = False) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise PreviewValidationError("Required numeric value is blank")
        return None
    if isinstance(value, bool):
        raise PreviewValidationError("Boolean is not a valid numeric value")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise PreviewValidationError(f"Invalid numeric value: {value!r}") from exc
    if not math.isfinite(result) or result < 0:
        raise PreviewValidationError(f"Numeric value must be finite and non-negative: {value!r}")
    return result


def _normalized_key(value: str | None) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", value).split()).casefold()


def _normalized_unit(value: Any) -> str | None:
    text = _text(value)
    return UnitNormalizer.normalize(text) if text is not None else None


def _same(left: Any, right: Any) -> bool:
    if left is None or right is None:
        return left is right
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=0, abs_tol=1e-9)
    return _normalized_key(str(left)) == _normalized_key(str(right))


def _validate_workbook_structure(workbook: Any) -> None:
    if tuple(workbook.sheetnames) != REQUIRED_SHEETS:
        raise PreviewValidationError("Workbook sheet names/order do not match the approved template")
    hidden = [name for name in REQUIRED_SHEETS if workbook[name].sheet_state != "visible"]
    if hidden:
        raise PreviewValidationError(f"Required sheet is not visible: {hidden[0]}")
    material_header = tuple(
        cell.value for cell in workbook["Материалын үнэ"][4][: len(MATERIAL_HEADERS)]
    )
    work_header = tuple(
        cell.value for cell in workbook["Ажил-Тээвэр-Механизм"][4][: len(WORK_HEADERS)]
    )
    if material_header != MATERIAL_HEADERS:
        raise PreviewValidationError("Material header row does not match the approved template")
    if work_header != WORK_HEADERS:
        raise PreviewValidationError("Work header row does not match the approved template")


def _nonempty_rows(sheet: Any, width: int) -> list[tuple[int, tuple[Any, ...]]]:
    rows: list[tuple[int, tuple[Any, ...]]] = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=5, max_col=width, values_only=True), start=5
    ):
        # Formula/template rows below the data range are not source records.
        # The approved template defines a record only when its external ID exists.
        if values[0] is not None and str(values[0]).strip():
            rows.append((row_number, tuple(values)))
    return rows


def _parse_work_rows(raw_rows: list[tuple[int, tuple[Any, ...]]], project_id: str) -> list[dict[str, Any]]:
    if len(raw_rows) != EXPECTED_WORK_ROWS:
        raise PreviewValidationError(f"Expected 65 work rows, found {len(raw_rows)}")
    parsed: list[dict[str, Any]] = []
    source_ids: list[str] = []
    canonical_ids: list[str] = []
    for row_number, row in raw_rows:
        source_id = _text(row[0], required=True)
        assert source_id is not None
        match = WORK_ID_PATTERN.fullmatch(source_id)
        if match is None:
            raise PreviewValidationError(f"Invalid Work ID at row {row_number}: {source_id}")
        canonical_id = f"{project_id}-WRK-{match.group(1)}"
        payload = {
            "work_id": canonical_id,
            "name": _text(row[2], required=True),
            "unit": _normalized_unit(row[4]),
            "quantity": _number(row[3]),
            "labor_unit_rate": _number(row[6]) if row[6] is not None else _number(row[5]),
            "status": "ACTIVE",
        }
        try:
            validated = ProjectWorkItemCreate.model_validate(payload)
        except ValidationError as exc:
            raise PreviewValidationError(f"Invalid work row {row_number}") from exc
        source_ids.append(source_id)
        canonical_ids.append(canonical_id)
        parsed.append({
            **validated.model_dump(mode="json"), "source_work_id": source_id,
            "source_row": row_number, "category": _text(row[1]),
        })
    if len(set(source_ids)) != EXPECTED_WORK_ROWS:
        raise PreviewValidationError("Duplicate Work ID detected")
    if len(set(canonical_ids)) != EXPECTED_WORK_ROWS:
        raise PreviewValidationError("Work ID mapping collision detected")
    return parsed


def _parse_material_rows(raw_rows: list[tuple[int, tuple[Any, ...]]]) -> list[dict[str, Any]]:
    if len(raw_rows) != EXPECTED_MATERIAL_ROWS:
        raise PreviewValidationError(f"Expected 986 material rows, found {len(raw_rows)}")
    parsed: list[dict[str, Any]] = []
    ids: list[str] = []
    current_price_count = 0
    new_price_count = 0
    prefixes = Counter()
    for row_number, row in raw_rows:
        material_id = _text(row[0], required=True)
        assert material_id is not None
        current_price = _number(row[5])
        new_price = _number(row[6])
        current_price_count += current_price is not None
        new_price_count += new_price is not None
        if material_id.startswith("PRD-ALTAI-B-"):
            prefixes["B"] += 1
        elif material_id.startswith("PRD-ALTAI-6R-"):
            prefixes["6R"] += 1
        else:
            raise PreviewValidationError(f"Unexpected Product ID prefix at row {row_number}")
        payload = {
            "material_id": material_id,
            "name": _text(row[2], required=True),
            "specification": _text(row[4]),
            "normalized_unit": _normalized_unit(row[3]),
            "unit_price": new_price if new_price is not None else current_price,
            "status": "ACTIVE",
        }
        try:
            validated = MaterialCreateRequest.model_validate(payload)
        except ValidationError as exc:
            raise PreviewValidationError(f"Invalid material row {row_number}") from exc
        ids.append(material_id)
        parsed.append({
            **validated.model_dump(mode="json"), "category": _text(row[1]),
            "current_price": current_price, "new_price": new_price,
            "source_row": row_number,
        })
    if len(set(ids)) != EXPECTED_MATERIAL_ROWS:
        raise PreviewValidationError("Duplicate Product ID detected")
    if prefixes != Counter({"B": 138, "6R": 848}):
        raise PreviewValidationError(f"Unexpected material prefix counts: {dict(prefixes)}")
    if current_price_count != 810 or new_price_count != 0:
        raise PreviewValidationError(
            f"Unexpected price counts: current={current_price_count}, new={new_price_count}"
        )
    return parsed


def _work_preview(parsed: Iterable[dict[str, Any]], snapshot: Any) -> list[WorkPreviewRow]:
    output: list[WorkPreviewRow] = []
    for row in parsed:
        existing = snapshot.work_items.get(row["work_id"])
        action = "CREATE"
        reason = ""
        if existing is not None:
            fields = ("name", "unit", "quantity", "labor_unit_rate", "status")
            differences = [field for field in fields if not _same(existing[field], row[field])]
            action = "SKIP" if not differences else "CONFLICT"
            reason = "" if not differences else f"Existing fields differ: {', '.join(differences)}"
        output.append(WorkPreviewRow(
            source_work_id=row["source_work_id"], canonical_work_id=row["work_id"],
            name=row["name"], unit=row["unit"], quantity=row["quantity"],
            labor_unit_rate=row["labor_unit_rate"], proposed_status="ACTIVE",
            action=action, conflict_reason=reason, source_row=row["source_row"],
            category=row["category"],
        ))
    return sorted(output, key=lambda item: item.canonical_work_id)


def _material_preview(parsed: list[dict[str, Any]], snapshot: Any) -> list[MaterialPreviewRow]:
    price_references: dict[tuple[str, str], set[float]] = defaultdict(set)
    semantic_groups: dict[tuple[str, str], list[str]] = defaultdict(list)
    for row in parsed:
        key = (_normalized_key(row["name"]), _normalized_key(row["normalized_unit"]))
        semantic_groups[key].append(row["material_id"])
        source_price = row["new_price"] if row["new_price"] is not None else row["current_price"]
        if source_price is not None:
            price_references[key].add(source_price)

    duplicate_keys = {
        key: f"SEM-{index:03d}"
        for index, key in enumerate(
            sorted((key for key, ids in semantic_groups.items() if len(ids) > 1)), start=1
        )
    }
    output: list[MaterialPreviewRow] = []
    for row in parsed:
        existing = snapshot.materials.get(row["material_id"])
        key = (_normalized_key(row["name"]), _normalized_key(row["normalized_unit"]))
        prices = sorted(price_references[key])
        source_price = row["new_price"] if row["new_price"] is not None else row["current_price"]
        source_kind = "excel_new" if row["new_price"] is not None else (
            "excel_current" if row["current_price"] is not None else "none"
        )
        existing_price = existing["unit_price"] if existing is not None else None
        candidate_price: float | None = None
        if source_price is not None:
            proposed_price = source_price
            resolution = source_kind
            status = "ACTIVE"
        elif existing_price is not None:
            proposed_price = float(existing_price)
            resolution = "db_price_preserved"
            status = "ACTIVE"
        elif len(prices) == 1:
            proposed_price = prices[0]
            resolution = "exact_reference_unique"
            status = "ACTIVE"
        elif len(prices) > 1:
            candidate_price = float(statistics.median(prices))
            proposed_price = None
            resolution = "exact_reference_ambiguous"
            status = "NEEDS_REVIEW"
        else:
            proposed_price = None
            resolution = "unresolved"
            status = "NEEDS_REVIEW"

        action = "CREATE" if status == "ACTIVE" else "REVIEW"
        reason = ""
        if existing is not None:
            comparable = {
                "name": row["name"], "specification": row["specification"],
                "normalized_unit": row["normalized_unit"], "unit_price": proposed_price,
                "status": status,
            }
            differences = [field for field, value in comparable.items() if not _same(existing[field], value)]
            action = "SKIP" if not differences else "CONFLICT"
            reason = "" if not differences else f"Existing fields differ: {', '.join(differences)}"
        output.append(MaterialPreviewRow(
            material_id=row["material_id"], name=row["name"],
            normalized_unit=row["normalized_unit"], specification=row["specification"],
            category=row["category"], source_price=source_price,
            source_price_kind=source_kind, existing_db_price=existing_price,
            proposed_unit_price=proposed_price, candidate_price=candidate_price,
            price_resolution=resolution, proposed_status=status,
            semantic_duplicate_group=duplicate_keys.get(key, ""), action=action,
            conflict_reason=reason, source_row=row["source_row"],
        ))
    return sorted(output, key=lambda item: item.material_id)


def _summary(
    workbook_path: Path,
    workbook_hash_before: str,
    database_path: Path,
    database_hash_before: str,
    project_id: str,
    work_rows: list[WorkPreviewRow],
    material_rows: list[MaterialPreviewRow],
    database_total_changes: int,
    generated_at_utc: str,
) -> dict[str, Any]:
    resolutions = Counter(row.price_resolution for row in material_rows)
    work_actions = Counter(row.action for row in work_rows)
    material_actions = Counter(row.action for row in material_rows)
    return {
        "schema_version": "1.0",
        "generated_at_utc": generated_at_utc,
        "source_workbook": {
            "path": str(workbook_path), "size": workbook_path.stat().st_size,
            "sha256_before": workbook_hash_before,
            "sha256_after": file_sha256(workbook_path),
        },
        "database": {
            "path": str(database_path), "sha256_before": database_hash_before,
            "sha256_after": file_sha256(database_path),
            "sqlite_total_changes": database_total_changes,
        },
        "project_id": project_id,
        "work_summary": {
            "source": len(work_rows), "valid": len(work_rows),
            "create": work_actions["CREATE"], "existing_identical": work_actions["SKIP"],
            "existing_conflict": work_actions["CONFLICT"],
        },
        "material_summary": {
            "source": len(material_rows), "valid": len(material_rows),
            "create": sum(row.action in {"CREATE", "REVIEW"} for row in material_rows),
            "existing_identical": material_actions["SKIP"],
            "existing_conflict": material_actions["CONFLICT"],
            "excel_price": resolutions["excel_new"] + resolutions["excel_current"],
            "db_price_preserved": resolutions["db_price_preserved"],
            "active_proposal": sum(row.proposed_status == "ACTIVE" for row in material_rows),
            "needs_review_proposal": sum(row.proposed_status == "NEEDS_REVIEW" for row in material_rows),
            "exact_reference_total": resolutions["exact_reference_unique"] + resolutions["exact_reference_ambiguous"],
            "exact_reference_unique": resolutions["exact_reference_unique"],
            "exact_reference_ambiguous": resolutions["exact_reference_ambiguous"],
            "unresolved": resolutions["unresolved"],
            "semantic_duplicate_groups": len({row.semantic_duplicate_group for row in material_rows if row.semantic_duplicate_group}),
        },
        "warnings": [
            "NEEDS_REVIEW prices are provisional or unresolved and require later research and approval.",
            "This report is a preview and does not import data.",
        ],
        "excluded_scopes": ["equipment", "transport", "work_material_links"],
        "database_writes_performed": False,
    }


def run_preview(
    *,
    file_path: Path,
    project_id: str,
    database_path: Path,
    report_dir: Path,
    repository_root: Path | None = None,
    generated_at_utc: str | None = None,
) -> PreviewResult:
    """Validate, compare, and report without changing workbook or database."""
    workbook_path = file_path.expanduser().resolve(strict=True)
    db_path = database_path.expanduser().resolve(strict=True)
    if workbook_path.suffix.lower() != ".xlsx":
        raise PreviewValidationError("--file must be an .xlsx workbook")
    project = _text(project_id, required=True)
    assert project is not None
    root = (repository_root or Path(__file__).resolve().parents[2]).resolve()
    output_dir = validate_report_directory(report_dir, root, workbook_path, db_path)
    workbook_hash_before = file_sha256(workbook_path)
    database_hash_before = file_sha256(db_path)

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        _validate_workbook_structure(workbook)
        material_raw = _nonempty_rows(workbook["Материалын үнэ"], len(MATERIAL_HEADERS))
        work_raw = _nonempty_rows(workbook["Ажил-Тээвэр-Механизм"], len(WORK_HEADERS))
        materials = _parse_material_rows(material_raw)
        works = _parse_work_rows(work_raw, project)
    finally:
        workbook.close()

    snapshot = load_database_snapshot(db_path, project)
    work_rows = _work_preview(works, snapshot)
    material_rows = _material_preview(materials, snapshot)
    timestamp = generated_at_utc or datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    summary = _summary(
        workbook_path, workbook_hash_before, db_path, database_hash_before,
        project, work_rows, material_rows, snapshot.total_changes, timestamp,
    )
    if summary["source_workbook"]["sha256_after"] != workbook_hash_before:
        raise PreviewValidationError("Workbook changed during preview")
    if summary["database"]["sha256_after"] != database_hash_before:
        raise PreviewValidationError("Database changed during preview")
    report_paths = write_preview_reports(output_dir, summary, work_rows, material_rows)
    return PreviewResult(summary, work_rows, material_rows, report_paths)
