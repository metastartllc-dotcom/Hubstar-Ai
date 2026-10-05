"""Read-only preview for canonical Work Master material recipes."""

from __future__ import annotations

import hashlib
import math
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.exporters.work_master_material_recipe_preview_reports import (
    REPORT_FILENAMES,
    write_recipe_preview_reports,
)
from app.repositories.work_master_material_recipe_preview import load_recipe_database_snapshot
from app.schemas.work_master_material_recipe_preview import (
    AliasPreviewRow,
    RecipePreviewResult,
    RecipePreviewRow,
)

EXPECTED_WORK_MASTERS = 106
EXPECTED_MATERIALS = 867
EXPECTED_RECIPES = 733
EXPECTED_ACTIVE = 654
EXPECTED_REVIEW = 79
EXPECTED_MATERIAL_ALIASES = 165
PROTECTED_WORK_MASTER_ID = "WKM-000001"

WORK_HEADERS = (
    "work_master_id", "name", "category", "default_unit",
    "default_labor_unit_rate", "status", "source_dataset",
    "source_work_id", "base_work_qty", "quantity_basis",
)
MATERIAL_HEADERS = (
    "material_id", "master_id", "code", "name", "specification",
    "normalized_unit", "unit_price", "status", "source", "source_count",
)
RECIPE_HEADERS = (
    "work_master_id", "material_id", "consumption_rate",
    "waste_percentage", "norm_with_waste", "source_material_qty",
    "base_work_qty", "status", "source_dataset", "source_rows",
    "source_link_count",
)
ALIAS_HEADERS = ("entity_type", "alias_id", "canonical_id", "reason")


class RecipePreviewValidationError(ValueError):
    """Raised when preview inputs violate the approved contract."""


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def validate_report_directory(
    report_dir: Path,
    repository_root: Path,
    workbook_path: Path,
    database_path: Path,
) -> Path:
    resolved = report_dir.resolve()
    forbidden = (
        repository_root.resolve(),
        workbook_path.resolve(strict=True).parent,
        database_path.resolve(strict=True).parent,
    )
    if any(_is_relative_to(resolved, parent) for parent in forbidden):
        raise RecipePreviewValidationError(
            "Report directory must be outside repository, workbook, and database directories"
        )
    if resolved.exists() and (
        not resolved.is_dir() or any(resolved.iterdir())
    ):
        raise RecipePreviewValidationError("Report directory must be a new or empty directory")
    if any((resolved / name).exists() for name in REPORT_FILENAMES):
        raise RecipePreviewValidationError("Recipe preview report already exists")
    return resolved


def _text(value: Any, *, required: bool = False) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise RecipePreviewValidationError("Required text value is blank")
        return None
    if isinstance(value, bool):
        raise RecipePreviewValidationError("Boolean is not valid text")
    return " ".join(unicodedata.normalize("NFKC", str(value)).split())


def _number(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise RecipePreviewValidationError("Boolean is not valid numeric data")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise RecipePreviewValidationError(f"Invalid numeric value: {value!r}") from exc


def _optional_non_negative(value: Any, field: str) -> float | None:
    result = _number(value)
    if result is not None and (not math.isfinite(result) or result < 0):
        raise RecipePreviewValidationError(f"{field} must be finite and non-negative")
    return result


def _records(workbook: Any, sheet_name: str, headers: tuple[str, ...]) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise RecipePreviewValidationError(f"Required sheet is missing: {sheet_name}")
    sheet = workbook[sheet_name]
    if tuple(cell.value for cell in sheet[1][: len(headers)]) != headers:
        raise RecipePreviewValidationError(f"Unexpected header contract: {sheet_name}")
    result = []
    for source_row, values in enumerate(
        sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True), start=2
    ):
        if any(value is not None for value in values):
            result.append({**dict(zip(headers, values)), "_source_row": source_row})
    return result


def _duplicates(values: list[Any]) -> set[Any]:
    return {value for value, count in Counter(values).items() if count > 1}


def _recipe_conflicts(
    work_id: str,
    material_id: str,
    status: str,
    rate: float | None,
    waste: float | None,
    work_unit: str | None,
    material_unit: str | None,
    duplicate_pairs: set[tuple[str, str]],
    work_ids: set[str],
    material_ids: set[str],
    alias_ids: set[str],
    source_link_count: float | None,
) -> list[str]:
    conflicts = []
    if (work_id, material_id) in duplicate_pairs:
        conflicts.append("DUPLICATE_WORK_MATERIAL_PAIR")
    if work_id not in work_ids:
        conflicts.append("MISSING_WORK_MASTER_REFERENCE")
    if material_id not in material_ids:
        conflicts.append("MISSING_MATERIAL_REFERENCE")
    if material_id in alias_ids:
        conflicts.append("MATERIAL_ALIAS_USED_INSTEAD_OF_CANONICAL_ID")
    if status == "ACTIVE":
        if rate is None or not math.isfinite(rate) or rate <= 0:
            conflicts.append("ACTIVE_CONSUMPTION_RATE_MUST_BE_FINITE_AND_POSITIVE")
        if work_unit is None:
            conflicts.append("ACTIVE_WORK_UNIT_MISSING")
        if material_unit is None:
            conflicts.append("ACTIVE_MATERIAL_UNIT_MISSING")
    elif status == "NEEDS_REVIEW":
        if rate is not None and (not math.isfinite(rate) or rate <= 0):
            conflicts.append("CONSUMPTION_RATE_MUST_BE_FINITE_AND_POSITIVE_WHEN_PRESENT")
    else:
        conflicts.append("UNSUPPORTED_RECIPE_STATUS")
    if waste is None or not math.isfinite(waste) or not 0 <= waste <= 100:
        conflicts.append("WASTE_PERCENTAGE_MUST_BE_BETWEEN_0_AND_100")
    if source_link_count is None or not math.isfinite(source_link_count) or source_link_count < 1:
        conflicts.append("SOURCE_LINK_COUNT_MUST_BE_POSITIVE")
    return conflicts


def run_recipe_preview(
    file_path: Path,
    project_id: str,
    database_path: Path,
    report_dir: Path,
    repository_root: Path | None = None,
) -> RecipePreviewResult:
    workbook_path = file_path.resolve(strict=True)
    database = database_path.resolve(strict=True)
    repository = (repository_root or Path.cwd()).resolve(strict=True)
    reports = validate_report_directory(report_dir, repository, workbook_path, database)
    workbook_hash = file_sha256(workbook_path)
    database_hash = file_sha256(database)
    snapshot = load_recipe_database_snapshot(database, project_id)
    if snapshot.total_changes != 0:
        raise RecipePreviewValidationError("Read-only database snapshot reported changes")
    if len(snapshot.protected_links) != 6:
        raise RecipePreviewValidationError("Expected exactly 6 protected facade links")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        works = _records(workbook, "work_masters", WORK_HEADERS)
        materials = _records(workbook, "materials", MATERIAL_HEADERS)
        recipes = _records(workbook, "work_material_norms", RECIPE_HEADERS)
        aliases = _records(workbook, "id_aliases", ALIAS_HEADERS)
    finally:
        workbook.close()
    expected = (
        (len(works), EXPECTED_WORK_MASTERS, "Work Masters"),
        (len(materials), EXPECTED_MATERIALS, "Materials"),
        (len(recipes), EXPECTED_RECIPES, "recipes"),
    )
    for actual, wanted, label in expected:
        if actual != wanted:
            raise RecipePreviewValidationError(f"Expected {wanted} {label}, found {actual}")

    work_ids = [_text(row["work_master_id"], required=True) for row in works]
    material_ids = [_text(row["material_id"], required=True) for row in materials]
    if _duplicates(work_ids):
        raise RecipePreviewValidationError("Duplicate Work Master ID detected")
    if _duplicates(material_ids):
        raise RecipePreviewValidationError("Duplicate Material ID detected")
    work_by_id = {
        str(identifier): {"name": _text(row["name"], required=True), "unit": _text(row["default_unit"])}
        for identifier, row in zip(work_ids, works)
    }
    material_by_id = {
        str(identifier): {"name": _text(row["name"], required=True), "unit": _text(row["normalized_unit"])}
        for identifier, row in zip(material_ids, materials)
    }

    material_aliases = [row for row in aliases if _text(row["entity_type"], required=True) == "MATERIAL"]
    if len(material_aliases) != EXPECTED_MATERIAL_ALIASES:
        raise RecipePreviewValidationError(
            f"Expected {EXPECTED_MATERIAL_ALIASES} material aliases, found {len(material_aliases)}"
        )
    alias_ids = [_text(row["alias_id"], required=True) for row in material_aliases]
    duplicate_aliases = _duplicates(alias_ids)
    alias_id_set = {str(value) for value in alias_ids}
    alias_rows = []
    for row in material_aliases:
        alias_id = _text(row["alias_id"], required=True) or ""
        canonical_id = _text(row["canonical_id"], required=True) or ""
        conflicts = []
        if alias_id in duplicate_aliases:
            conflicts.append("DUPLICATE_ALIAS_ID")
        if alias_id == canonical_id:
            conflicts.append("ALIAS_EQUALS_CANONICAL_ID")
        if alias_id in material_by_id:
            conflicts.append("ALIAS_COLLIDES_WITH_CANONICAL_MATERIAL_ID")
        if canonical_id not in material_by_id:
            conflicts.append("MISSING_CANONICAL_MATERIAL")
        if canonical_id in alias_id_set:
            conflicts.append("ALIAS_CHAIN_OR_CYCLE_NOT_ALLOWED")
        alias_rows.append(AliasPreviewRow(
            alias_id=alias_id,
            canonical_id=canonical_id,
            reason=_text(row["reason"]) or "",
            validation_status="INVALID" if conflicts else "VALID",
            conflict_reason=";".join(conflicts),
        ))

    pairs = [
        (_text(row["work_master_id"], required=True) or "", _text(row["material_id"], required=True) or "")
        for row in recipes
    ]
    duplicate_pairs = _duplicates(pairs)
    recipe_rows = []
    for row, (work_id, material_id) in zip(recipes, pairs):
        status = _text(row["status"], required=True) or ""
        rate = _number(row["consumption_rate"])
        waste = _number(row["waste_percentage"])
        source_link_count = _number(row["source_link_count"])
        work = work_by_id.get(work_id, {"name": "", "unit": None})
        material = material_by_id.get(material_id, {"name": "", "unit": None})
        conflicts = _recipe_conflicts(
            work_id, material_id, status, rate, waste, work["unit"], material["unit"],
            duplicate_pairs, set(work_by_id), set(material_by_id), alias_id_set,
            source_link_count,
        )
        missing = []
        if work_id not in snapshot.work_master_ids:
            missing.append("WORK_MASTER")
        if material_id not in snapshot.material_ids:
            missing.append("MATERIAL")
        reference_status = "PRODUCTION_NOT_READY" if missing else "READY"
        readiness = (
            "PRODUCTION_NOT_READY"
            if missing or conflicts or status != "ACTIVE"
            else "READY"
        )
        action = "INVALID" if conflicts else (
            "ACTIVE_PROPOSAL" if status == "ACTIVE" else "EXCLUDED_REVIEW"
        )
        recipe_rows.append(RecipePreviewRow(
            work_master_id=work_id,
            material_id=material_id,
            work_name=str(work["name"]),
            material_name=str(material["name"]),
            work_unit=work["unit"],
            material_unit=material["unit"],
            consumption_rate=rate,
            waste_percentage=waste,
            norm_with_waste=_optional_non_negative(row["norm_with_waste"], "norm_with_waste"),
            source_material_qty=_optional_non_negative(row["source_material_qty"], "source_material_qty"),
            base_work_qty=_optional_non_negative(row["base_work_qty"], "base_work_qty"),
            source_status=status,
            action=action,
            internal_conflict_reason=";".join(conflicts),
            production_reference_status=reference_status,
            production_readiness=readiness,
            production_readiness_reason=";".join(filter(None, (
                "INTERNAL_VALIDATION_FAILED" if conflicts else "",
                "REVIEW_ONLY" if status == "NEEDS_REVIEW" else "",
                "" if not missing else "MISSING_PRODUCTION_" + "_AND_".join(missing),
            ))),
            source_dataset=_text(row["source_dataset"], required=True) or "",
            source_rows=_text(row["source_rows"], required=True) or "",
            source_link_count=int(source_link_count or 0),
        ))

    status_counts = Counter(row.source_status for row in recipe_rows)
    if status_counts != Counter({"ACTIVE": EXPECTED_ACTIVE, "NEEDS_REVIEW": EXPECTED_REVIEW}):
        raise RecipePreviewValidationError(f"Unexpected recipe status counts: {dict(status_counts)}")
    internal_recipe_conflicts = sum(bool(row.internal_conflict_reason) for row in recipe_rows)
    invalid_aliases = sum(row.validation_status == "INVALID" for row in alias_rows)
    production_reference_counts = Counter(
        row.production_reference_status for row in recipe_rows
    )
    production_counts = Counter(row.production_readiness for row in recipe_rows)
    active_counts = Counter(
        row.production_readiness for row in recipe_rows if row.source_status == "ACTIVE"
    )
    active_reference_counts = Counter(
        row.production_reference_status
        for row in recipe_rows
        if row.source_status == "ACTIVE"
    )
    summary = {
        "counts": {
            "work_masters": len(works),
            "materials": len(materials),
            "recipes": len(recipe_rows),
            "source_active": status_counts["ACTIVE"],
            "source_needs_review": status_counts["NEEDS_REVIEW"],
            "active_valid": sum(
                row.source_status == "ACTIVE" and not row.internal_conflict_reason
                for row in recipe_rows
            ),
            "excluded_review": status_counts["NEEDS_REVIEW"],
            "material_aliases": len(alias_rows),
            "internal_recipe_conflicts": internal_recipe_conflicts,
            "invalid_aliases": invalid_aliases,
            "internal_conflicts_total": internal_recipe_conflicts + invalid_aliases,
            "protected_project_links": len(snapshot.protected_links),
            "importable_active": sum(
                row.action == "ACTIVE_PROPOSAL"
                and row.production_readiness == "READY"
                for row in recipe_rows
            ),
        },
        "database": {
            "filename": database.name,
            "revision": snapshot.revision,
            "sha256": database_hash,
            "sqlite_total_changes": snapshot.total_changes,
        },
        "operation_mode": "DRY_RUN",
        "production_readiness": {
            "all_rows": {
                "ready": production_counts["READY"],
                "production_not_ready": production_counts["PRODUCTION_NOT_READY"],
            },
            "active_rows": {
                "ready": active_counts["READY"],
                "production_not_ready": active_counts["PRODUCTION_NOT_READY"],
            },
        },
        "production_reference_readiness": {
            "all_rows": {
                "ready": production_reference_counts["READY"],
                "production_not_ready": production_reference_counts["PRODUCTION_NOT_READY"],
            },
            "active_rows": {
                "ready": active_reference_counts["READY"],
                "production_not_ready": active_reference_counts["PRODUCTION_NOT_READY"],
            },
        },
        "reconciliation": {
            "recipes_equal_active_plus_review": (
                len(recipe_rows) == status_counts["ACTIVE"] + status_counts["NEEDS_REVIEW"]
            ),
            "active_equal_valid_plus_invalid": (
                status_counts["ACTIVE"]
                == sum(
                    row.source_status == "ACTIVE" and not row.internal_conflict_reason
                    for row in recipe_rows
                )
                + sum(
                    row.source_status == "ACTIVE" and bool(row.internal_conflict_reason)
                    for row in recipe_rows
                )
            ),
            "references_equal_ready_plus_not_ready": (
                len(recipe_rows)
                == production_reference_counts["READY"]
                + production_reference_counts["PRODUCTION_NOT_READY"]
            ),
        },
        "source_policy": {
            "needs_review_imported": False,
            "proposed_name_or_unit_applied": False,
            "unit_conversion_applied": False,
        },
        "status": "BLOCKED_INTERNAL_CONFLICTS" if internal_recipe_conflicts or invalid_aliases else "PREVIEW_READY",
        "workbook": {"filename": workbook_path.name, "sha256": workbook_hash},
    }
    report_paths = write_recipe_preview_reports(
        reports, summary, recipe_rows, alias_rows, list(snapshot.protected_links)
    )
    if file_sha256(workbook_path) != workbook_hash:
        raise RecipePreviewValidationError("Source workbook changed during preview")
    if file_sha256(database) != database_hash:
        raise RecipePreviewValidationError("Database changed during preview")
    final_snapshot = load_recipe_database_snapshot(database, project_id)
    if final_snapshot.total_changes != 0:
        raise RecipePreviewValidationError("Preview database reported writes")
    return RecipePreviewResult(
        summary=summary,
        recipe_rows=sorted(recipe_rows, key=lambda row: (row.work_master_id, row.material_id)),
        alias_rows=sorted(alias_rows, key=lambda row: row.alias_id),
        protected_links=list(snapshot.protected_links),
        report_paths=report_paths,
    )
