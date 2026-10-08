"""Canonical workbook versus production master-data reconciliation preview."""

from __future__ import annotations

import json
import math
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.exporters.canonical_master_data_reconciliation_reports import (
    REPORT_FILENAMES, write_reconciliation_reports,
)
from app.importers.work_master_material_recipe_preview import (
    ALIAS_HEADERS, MATERIAL_HEADERS, RECIPE_HEADERS, WORK_HEADERS,
    file_sha256, validate_report_directory,
)
from app.repositories.canonical_master_data_reconciliation import load_production_snapshot
from app.schemas.canonical_master_data_reconciliation import (
    AliasReconciliationRow, EntityReconciliationRow, ReconciliationResult,
    RecipeReadinessRow,
)

WORK_FIELDS = ("name", "category", "default_unit", "default_labor_unit_rate", "status", "source_dataset", "source_work_id")
MATERIAL_FIELDS = ("master_id", "code", "name", "specification", "normalized_unit", "unit_price", "status")


class ReconciliationValidationError(ValueError):
    """Raised when reconciliation cannot run safely."""


def _records(workbook: Any, sheet_name: str, headers: tuple[str, ...]) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise ReconciliationValidationError(f"Required sheet is missing: {sheet_name}")
    sheet = workbook[sheet_name]
    if tuple(cell.value for cell in sheet[1][:len(headers)]) != headers:
        raise ReconciliationValidationError(f"Unexpected header contract: {sheet_name}")
    return [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True) if any(value is not None for value in values)]


def _value(value: Any) -> Any:
    if hasattr(value, "value"):
        value = value.value
    if isinstance(value, str):
        value = " ".join(value.split())
        return value or None
    return value


def _diffs(workbook: dict[str, Any], production: dict[str, Any], fields: tuple[str, ...], preserve_price: bool = False) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for field in fields:
        left, right = _value(workbook.get(field)), _value(production.get(field))
        if left != right:
            resolution = "REVIEW_REQUIRED"
            if preserve_price and field == "unit_price" and left is None and right is not None:
                resolution = "PRESERVE_PRODUCTION_NON_NULL_PRICE"
            result[field] = {"workbook": left, "production": right, "resolution": resolution}
    return result


def _entity_rows(
    entity_type: str, workbook: dict[str, dict[str, Any]], production: dict[str, dict[str, Any]],
    fields: tuple[str, ...], source_fields: tuple[str, ...] = (), preserve_price: bool = False,
) -> list[EntityReconciliationRow]:
    source_owner = {
        tuple(_value(row.get(field)) for field in source_fields): external_id
        for external_id, row in production.items()
        if source_fields and all(_value(row.get(field)) is not None for field in source_fields)
    }
    workbook_source_counts = Counter(
        tuple(_value(row.get(field)) for field in source_fields)
        for row in workbook.values()
        if source_fields and all(_value(row.get(field)) is not None for field in source_fields)
    )
    rows = []
    for external_id in sorted(set(workbook) | set(production)):
        wb, db = workbook.get(external_id), production.get(external_id)
        source_identity = ""
        action, diffs, policy = "", {}, ""
        if wb is not None and source_fields:
            pair = tuple(_value(wb.get(field)) for field in source_fields)
            source_identity = "|".join("" if value is None else str(value) for value in pair)
            owner = source_owner.get(pair) if all(value is not None for value in pair) else None
            if workbook_source_counts[pair] > 1:
                action, diffs, policy = "CONFLICT", {"workbook_source_identity_count": workbook_source_counts[pair]}, "FAIL_CLOSED_SOURCE_COLLISION"
            elif owner is not None and owner != external_id:
                action, diffs, policy = "CONFLICT", {"source_identity_owner": owner}, "FAIL_CLOSED_SOURCE_COLLISION"
        if not action:
            if wb is None:
                action, policy = "RETAIN_PRODUCTION_ONLY", "NO_DELETE_OR_MERGE"
            elif db is None:
                action, policy = "CREATE_PROPOSAL", "PROPOSAL_ONLY_NO_WRITE"
            else:
                diffs = _diffs(wb, db, fields, preserve_price)
                action = "UPDATE_REVIEW_REQUIRED" if diffs else "KEEP_IDENTICAL"
                policy = "EXPLICIT_REVIEW_NO_AUTOMATIC_SELECTION" if diffs else "NO_CHANGE"
        rows.append(EntityReconciliationRow(
            entity_type, external_id, action, wb is not None, db is not None,
            source_identity, json.dumps(diffs, ensure_ascii=False, sort_keys=True, separators=(",", ":")), policy,
        ))
    return rows


def _alias_rows(aliases: list[dict[str, Any]], works: dict[str, Any], materials: dict[str, Any], production_materials: dict[str, Any]) -> list[AliasReconciliationRow]:
    alias_ids = {str(_value(row["alias_id"])) for row in aliases}
    result = []
    for row in aliases:
        entity = str(_value(row["entity_type"]))
        alias_id, target = str(_value(row["alias_id"])), str(_value(row["canonical_id"]))
        issues = []
        interpretation = "MATERIAL_ALIAS_REVIEW_ONLY" if entity == "MATERIAL" else "WORK_PACKAGE_MAPPING_NOT_WORK_MASTER_ALIAS"
        targets = materials if entity == "MATERIAL" else works
        if target not in targets:
            issues.append("MISSING_CANONICAL_TARGET")
        if alias_id == target:
            issues.append("ALIAS_EQUALS_TARGET")
        if target in alias_ids:
            issues.append("ALIAS_CHAIN_OR_CYCLE")
        if entity == "MATERIAL" and alias_id in materials:
            issues.append("ALIAS_COLLIDES_WITH_CANONICAL_ID")
        if entity == "MATERIAL" and alias_id in production_materials:
            interpretation += ";PRODUCTION_ROW_MUST_BE_RETAINED"
        action = "CONFLICT" if issues else "ALIAS_REVIEW_REQUIRED"
        result.append(AliasReconciliationRow(entity, alias_id, target, action, str(_value(row.get("reason")) or ""), ";".join(issues), interpretation))
    return sorted(result, key=lambda row: (row.entity_type, row.alias_id, row.canonical_id))


def _recipe_rows(recipes: list[dict[str, Any]], works: dict[str, dict[str, Any]], materials: dict[str, dict[str, Any]], snapshot: Any, work_actions: dict[str, str], material_actions: dict[str, str], alias_targets: set[str]) -> list[RecipeReadinessRow]:
    pairs = [(str(_value(row["work_master_id"])), str(_value(row["material_id"]))) for row in recipes]
    duplicates = {pair for pair, count in Counter(pairs).items() if count > 1}
    result = []
    for row, (work_id, material_id) in zip(recipes, pairs):
        status = str(_value(row["status"]))
        rate, waste = row.get("consumption_rate"), row.get("waste_percentage")
        try:
            rate_number = None if rate is None or isinstance(rate, bool) else float(rate)
        except (TypeError, ValueError):
            rate_number = None
        try:
            waste_number = None if waste is None or isinstance(waste, bool) else float(waste)
        except (TypeError, ValueError):
            waste_number = None
        issues = []
        if (work_id, material_id) in duplicates:
            issues.append("DUPLICATE_PAIR")
        if work_id not in works:
            issues.append("MISSING_WORKBOOK_WORK")
        if material_id not in materials:
            issues.append("MISSING_WORKBOOK_MATERIAL")
        work_unit = _value(works.get(work_id, {}).get("default_unit"))
        material_unit = _value(materials.get(material_id, {}).get("normalized_unit"))
        if status == "ACTIVE":
            if rate_number is None or not math.isfinite(rate_number) or rate_number <= 0:
                issues.append("INVALID_ACTIVE_RATE")
            if work_unit is None:
                issues.append("MISSING_WORK_UNIT")
            if material_unit is None:
                issues.append("MISSING_MATERIAL_UNIT")
        if waste_number is None or not math.isfinite(waste_number) or not 0 <= waste_number <= 100:
            issues.append("INVALID_WASTE")
        workbook_work_ref, workbook_material_ref = work_id in works, material_id in materials
        work_ref, material_ref = work_id in snapshot.work_masters, material_id in snapshot.materials
        semantic_issue = (
            work_ref and workbook_work_ref
            and _value(works[work_id].get("name")) != _value(snapshot.work_masters[work_id].get("name"))
        ) or (
            material_ref and workbook_material_ref
            and (
                _value(materials[material_id].get("name")) != _value(snapshot.materials[material_id].get("name"))
                or _value(materials[material_id].get("specification")) != _value(snapshot.materials[material_id].get("specification"))
            )
        )
        if semantic_issue:
            semantic_match = "MISMATCH"
        elif work_ref and material_ref and workbook_work_ref and workbook_material_ref:
            semantic_match = "MATCH"
        else:
            semantic_match = "NOT_EVALUATED"
        unit_issues: list[str] = []
        comparison_issues: list[str] = []
        if work_unit is None:
            unit_issues.append("WORK_UNIT_MISSING")
            if work_ref and _value(snapshot.work_masters[work_id].get("default_unit")) is not None:
                comparison_issues.append("WORK_UNIT_MISSING")
        elif work_ref and work_unit != _value(snapshot.work_masters[work_id].get("default_unit")):
            unit_issues.append("WORK_UNIT_MISMATCH")
            comparison_issues.append("WORK_UNIT_MISMATCH")
        if material_unit is None:
            unit_issues.append("MATERIAL_UNIT_MISSING")
            if material_ref and _value(snapshot.materials[material_id].get("normalized_unit")) is not None:
                comparison_issues.append("MATERIAL_UNIT_MISSING")
        elif material_ref and material_unit != _value(snapshot.materials[material_id].get("normalized_unit")):
            unit_issues.append("MATERIAL_UNIT_MISMATCH")
            comparison_issues.append("MATERIAL_UNIT_MISMATCH")
        if any(issue.endswith("_MISSING") for issue in comparison_issues):
            unit_match = "MISSING"
        elif comparison_issues:
            unit_match = "MISMATCH"
        elif work_ref and material_ref:
            unit_match = "MATCH"
        else:
            unit_match = "NOT_EVALUATED"
        dependencies: set[str] = set()
        actions = {work_actions.get(work_id), material_actions.get(material_id)}
        for action in ("CONFLICT", "CREATE_PROPOSAL", "UPDATE_REVIEW_REQUIRED"):
            if action in actions:
                dependencies.add(action)
        if material_id in alias_targets:
            dependencies.add("ALIAS_REVIEW_REQUIRED")
        if semantic_match == "MISMATCH":
            dependencies.add("SEMANTIC_MISMATCH")
        if unit_issues:
            dependencies.add("UNIT_REVIEW_REQUIRED")
        if issues:
            readiness = "NOT_IMPORTABLE_INVALID"
        elif status == "NEEDS_REVIEW":
            readiness = "NOT_IMPORTABLE_REVIEW"
        elif "CONFLICT" in dependencies:
            readiness = "BLOCKED_CONFLICT"
        elif "CREATE_PROPOSAL" in dependencies:
            readiness = "DEPENDS_ON_CREATE"
        elif "UPDATE_REVIEW_REQUIRED" in dependencies or "SEMANTIC_MISMATCH" in dependencies or "UNIT_REVIEW_REQUIRED" in dependencies:
            readiness = "DEPENDS_ON_UPDATE_REVIEW"
        elif "ALIAS_REVIEW_REQUIRED" in dependencies:
            readiness = "DEPENDS_ON_ALIAS_REVIEW"
        elif work_ref and material_ref:
            readiness = "READY_EXISTING"
        else:
            readiness = "PRODUCTION_NOT_READY"
        result.append(RecipeReadinessRow(
            work_id, material_id, status, "INVALID" if issues else "VALID", ";".join(issues),
            workbook_work_ref, workbook_material_ref, work_ref, material_ref,
            semantic_match, unit_match, ";".join(unit_issues),
            ";".join(sorted(dependencies)) or "NONE", readiness,
            readiness == "READY_EXISTING",
        ))
    return sorted(result, key=lambda item: (item.work_master_id, item.material_id))


def run_reconciliation_preview(file_path: Path, project_id: str, database_path: Path, report_dir: Path, repository_root: Path | None = None) -> ReconciliationResult:
    workbook_path, database = file_path.resolve(strict=True), database_path.resolve(strict=True)
    repository = (repository_root or Path.cwd()).resolve(strict=True)
    reports = validate_report_directory(report_dir, repository, workbook_path, database)
    workbook_hash, database_hash = file_sha256(workbook_path), file_sha256(database)
    snapshot = load_production_snapshot(database, project_id)
    if snapshot.revision != "0004_work_masters" or snapshot.total_changes != 0:
        raise ReconciliationValidationError("Production database is not ready for read-only reconciliation")
    book = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        work_records = _records(book, "work_masters", WORK_HEADERS)
        material_records = _records(book, "materials", MATERIAL_HEADERS)
        recipe_records = _records(book, "work_material_norms", RECIPE_HEADERS)
        aliases = _records(book, "id_aliases", ALIAS_HEADERS)
    finally:
        book.close()
    if (len(work_records), len(material_records), len(recipe_records)) != (106, 867, 733):
        raise ReconciliationValidationError("Canonical workbook counts do not match 106/867/733")
    works = {str(_value(row["work_master_id"])): row for row in work_records}
    materials = {str(_value(row["material_id"])): row for row in material_records}
    if len(works) != 106 or len(materials) != 867:
        raise ReconciliationValidationError("Duplicate canonical external ID")
    work_rows = _entity_rows("WORK_MASTER", works, snapshot.work_masters, WORK_FIELDS, ("source_dataset", "source_work_id"))
    material_rows = _entity_rows("MATERIAL", materials, snapshot.materials, MATERIAL_FIELDS, preserve_price=True)
    alias_rows = _alias_rows(aliases, works, materials, snapshot.materials)
    work_actions = {row.external_id: row.action for row in work_rows}
    material_actions = {row.external_id: row.action for row in material_rows}
    alias_targets = {
        row.canonical_id for row in alias_rows
        if row.entity_type == "MATERIAL" and row.action == "ALIAS_REVIEW_REQUIRED"
    }
    recipe_rows = _recipe_rows(recipe_records, works, materials, snapshot, work_actions, material_actions, alias_targets)
    protected_rows = []
    for protected in snapshot.protected_rows:
        if protected.entity_type != "FACADE_MATERIAL_LINK":
            protected_rows.append(protected)
            continue
        data = json.loads(protected.snapshot_json)
        material_id = protected.external_id
        workbook_material = materials.get(material_id, {})
        production_material = snapshot.materials.get(material_id, {})
        workbook_price = _value(workbook_material.get("unit_price"))
        production_price = _value(production_material.get("unit_price"))
        workbook_unit = _value(workbook_material.get("normalized_unit"))
        production_unit = _value(production_material.get("normalized_unit"))
        data["reconciliation"] = {
            "price_raw_difference": workbook_price != production_price,
            "unit_raw_difference": workbook_unit != production_unit,
            "workbook_unit_price": workbook_price,
            "production_unit_price": production_price,
            "effective_unit_price": production_price if workbook_price is None and production_price is not None else None,
            "price_policy": "PRESERVE_PRODUCTION_NON_NULL_PRICE" if workbook_price is None and production_price is not None else "REVIEW_DIFFERENCE" if workbook_price != production_price else "NO_CHANGE",
            "workbook_unit": workbook_unit,
            "production_unit": production_unit,
            "unit_policy": "NO_CHANGE" if workbook_unit == production_unit else "REVIEW_REQUIRED_NO_CONVERSION",
        }
        protected_rows.append(type(protected)(protected.entity_type, protected.external_id, json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":")), protected.protection_policy))
    actions = Counter(row.action for row in work_rows + material_rows + alias_rows)
    readiness = Counter(row.readiness for row in recipe_rows)
    update_rows = [row for row in work_rows + material_rows if row.action == "UPDATE_REVIEW_REQUIRED"]
    field_entities: dict[str, set[str]] = {}
    field_occurrences: Counter[str] = Counter()
    field_names = {
        "default_unit": "unit", "normalized_unit": "unit", "unit_price": "price",
        "source_dataset": "source_identity", "source_work_id": "source_identity",
    }
    for field in ("name", "unit", "specification", "price", "status", "source_identity"):
        field_occurrences[field] = 0
        field_entities[field] = set()
    for row in update_rows:
        for field in json.loads(row.explicit_diffs):
            category = field_names.get(field, field)
            field_occurrences[category] += 1
            field_entities.setdefault(category, set()).add(f"{row.entity_type}:{row.external_id}")
    dependency_occurrences = Counter(
        dependency for row in recipe_rows
        for dependency in row.dependency_actions.split(";") if dependency != "NONE"
    )
    unit_breakdown = Counter((row.source_status, row.unit_match) for row in recipe_rows)
    unit_issue_breakdown = Counter(
        (row.source_status, issue)
        for row in recipe_rows for issue in row.unit_issues.split(";") if issue
    )
    conflicts = actions["CONFLICT"] + sum(row.internal_validation == "INVALID" for row in recipe_rows)
    summary = {
        "actions": dict(sorted(actions.items())),
        "counts": {
            "workbook_work_masters": len(works), "production_work_masters": len(snapshot.work_masters),
            "workbook_materials": len(materials), "production_materials": len(snapshot.materials),
            "material_aliases": sum(row.entity_type == "MATERIAL" for row in alias_rows),
            "work_package_mappings": sum(row.entity_type == "WORK_PACKAGE" for row in alias_rows),
            "recipes": len(recipe_rows), "blocking_conflicts": conflicts,
            "protected_rows": len(protected_rows),
        },
        "database": {"filename": database.name, "revision": snapshot.revision, "sha256": database_hash, "sqlite_total_changes": snapshot.total_changes},
        "operation_mode": "DRY_RUN", "recipe_readiness": dict(sorted(readiness.items())),
        "recipe_dependency_occurrences": dict(sorted(dependency_occurrences.items())),
        "recipe_unit_breakdown": {
            f"{status}:{unit_state}": count
            for (status, unit_state), count in sorted(unit_breakdown.items())
        },
        "recipe_unit_issue_breakdown": {
            f"{status}:{issue}": count
            for (status, issue), count in sorted(unit_issue_breakdown.items())
        },
        "update_review_breakdown": {
            "unique_entities": len(update_rows),
            "field_difference_occurrences": dict(sorted(field_occurrences.items())),
            "unique_entities_by_field": {
                field: len(entities) for field, entities in sorted(field_entities.items())
            },
        },
        "status": "BLOCKED_CONFLICTS" if conflicts else "PREVIEW_READY",
        "workbook": {"filename": workbook_path.name, "sha256": workbook_hash},
        "policies": {"unit_conversion": False, "proposed_name_unit_applied": False, "automatic_merge": False, "production_price_overwrite_by_blank": False},
    }
    paths = write_reconciliation_reports(reports, summary, work_rows, material_rows, alias_rows, recipe_rows, protected_rows)
    if file_sha256(workbook_path) != workbook_hash or file_sha256(database) != database_hash:
        raise ReconciliationValidationError("Workbook or database changed during reconciliation")
    if load_production_snapshot(database, project_id).total_changes != 0:
        raise ReconciliationValidationError("Read-only reconciliation reported database changes")
    return ReconciliationResult(summary, work_rows, material_rows, alias_rows, recipe_rows, protected_rows, paths)
